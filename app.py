import io
import os
import re
import shutil
import sqlite3
import uuid
from datetime import datetime, date
from functools import wraps

from flask import Flask, render_template, request, redirect, url_for, session, flash, g, send_file, send_from_directory, jsonify
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont, ImageOps

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "cambia-esta-clave-en-produccion")

BASE_DIR = os.path.dirname(__file__)

# En Render, /data es el disco persistente (sobrevive a los despliegues).
# Si no existe (ej. en una máquina local), se usa la carpeta del proyecto como antes.
DATA_DIR = "/data" if os.path.isdir("/data") else BASE_DIR
DB_PATH = os.path.join(DATA_DIR, "liga.db")
FOTOS_DIR = os.path.join(DATA_DIR, "fotos_jugadores")
os.makedirs(FOTOS_DIR, exist_ok=True)

# Migración única: si el disco persistente está recién montado y todavía no
# tiene datos, pero existen datos viejos en la carpeta efímera del proyecto
# (de antes de tener disco persistente), se copian una sola vez.
if DATA_DIR != BASE_DIR:
    _db_vieja = os.path.join(BASE_DIR, "liga.db")
    if not os.path.exists(DB_PATH) and os.path.exists(_db_vieja):
        shutil.copy2(_db_vieja, DB_PATH)

    _fotos_viejas = os.path.join(BASE_DIR, "static", "fotos_jugadores")
    if os.path.isdir(_fotos_viejas):
        for _nombre in os.listdir(_fotos_viejas):
            _destino = os.path.join(FOTOS_DIR, _nombre)
            if not os.path.exists(_destino):
                try:
                    shutil.copy2(os.path.join(_fotos_viejas, _nombre), _destino)
                except OSError:
                    pass

DATABASE_URL = os.environ.get("DATABASE_URL", "").strip()
USANDO_POSTGRES = bool(DATABASE_URL)

if USANDO_POSTGRES:
    import psycopg2
    import psycopg2.extras
    IntegrityError = psycopg2.IntegrityError
else:
    IntegrityError = sqlite3.IntegrityError


class DBWrapper:
    """Traduce placeholders '?' (estilo sqlite) a '%s' (estilo psycopg2)
    y adapta fetchone()/fetchall() para devolver dicts en ambos motores."""

    def __init__(self, conn):
        self.conn = conn

    def execute(self, query, params=()):
        cur = self.conn.cursor()
        if USANDO_POSTGRES:
            query = re.sub(r"\?", "%s", query)
            query = query.replace("INTEGER PRIMARY KEY AUTOINCREMENT", "SERIAL PRIMARY KEY")
        cur.execute(query, params)
        return CursorWrapper(cur)

    def commit(self):
        self.conn.commit()

    def rollback(self):
        self.conn.rollback()

    def close(self):
        self.conn.close()


class CursorWrapper:
    def __init__(self, cur):
        self.cur = cur

    def fetchone(self):
        row = self.cur.fetchone()
        return row

    def fetchall(self):
        return self.cur.fetchall()

    @property
    def lastrowid(self):
        if USANDO_POSTGRES:
            try:
                return self.cur.fetchone()[0]
            except Exception:
                return None
        return self.cur.lastrowid

ADMIN_USER = os.environ.get("ADMIN_USER", "admin")
ADMIN_PASS = os.environ.get("ADMIN_PASS", "Oyambarillo2026")
IMPRENTA_USER = os.environ.get("IMPRENTA_USER", "carnets")
IMPRENTA_PASS = os.environ.get("IMPRENTA_PASS", "")
CALIFICACION_USER = os.environ.get("CALIFICACION_USER", "comision")
CALIFICACION_PASS = os.environ.get("CALIFICACION_PASS", "")
SANCIONES_USER = os.environ.get("SANCIONES_USER", "sanciones")
SANCIONES_PASS = os.environ.get("SANCIONES_PASS", "")
TECNICA_USER = os.environ.get("TECNICA_USER", "tecnica")
TECNICA_PASS = os.environ.get("TECNICA_PASS", "")
VOCALIA_USER = os.environ.get("VOCALIA_USER", "vocalia")
VOCALIA_PASS = os.environ.get("VOCALIA_PASS", "")


@app.route("/fotos_jugadores/<path:filename>")
def fotos_jugadores(filename):
    return send_from_directory(FOTOS_DIR, filename)

CUPO_MAXIMO_EQUIPO = 35
CUPO_MAXIMO_JUVENIL = 3
CATEGORIA_ACTIVA = "Sub 45"
SUBCATEGORIAS = ["Sub 45", "Juvenil"]


def _categorias_liga(db):
    """Lista de nombres de categorías, en el orden en que se crearon.
    Editable por el admin desde /categorias -- no está fija en el código."""
    return [f["nombre"] for f in db.execute("SELECT nombre FROM categorias ORDER BY id").fetchall()]


def _categoria_valida(db, categoria):
    return categoria if categoria in _categorias_liga(db) else CATEGORIA_ACTIVA


def _edad_minima_categoria(db, categoria):
    fila = db.execute("SELECT edad_minima FROM categorias WHERE nombre = ?", (categoria,)).fetchone()
    return fila["edad_minima"] if fila else None


def _divisiones_de_categoria(db, categoria):
    """Nombres de las divisiones/series de una categoría (ej. Senior -> Máxima,
    Primera), en el orden en que se crearon. Lista vacía si esa categoría no
    se divide en series (ej. Sub 45)."""
    return [
        f["nombre"] for f in
        db.execute("SELECT nombre FROM divisiones WHERE categoria = ? ORDER BY id", (categoria,)).fetchall()
    ]


def _categoria_y_division(db, categoria_arg, division_arg):
    """Valida la categoría pedida y, si esa categoría tiene divisiones, la
    división pedida (o la primera por defecto). Si la categoría no tiene
    divisiones, la división siempre es '' (no aplica)."""
    categoria = _categoria_valida(db, categoria_arg)
    divisiones = _divisiones_de_categoria(db, categoria)
    if not divisiones:
        return categoria, ""
    division = division_arg if division_arg in divisiones else divisiones[0]
    return categoria, division

EQUIPOS_INICIALES = [
    "San Juan", "El Progreso", "La Union", "Santa Rosa",
    "Los Andes", "Independiente", "Deportivo Central", "Juventud",
]

# Formaciones disponibles para la alineación: cada slot tiene una posición en
# la cancha (top/left en % para ubicarlo con CSS) y una etiqueta corta.
FORMACIONES = {
    "4-3-3": [
        {"code": "GK", "label": "POR", "top": 90, "left": 50},
        {"code": "LB", "label": "LI", "top": 72, "left": 16},
        {"code": "CB1", "label": "DFC", "top": 76, "left": 36},
        {"code": "CB2", "label": "DFC", "top": 76, "left": 64},
        {"code": "RB", "label": "LD", "top": 72, "left": 84},
        {"code": "CDM1", "label": "MCD", "top": 54, "left": 36},
        {"code": "CDM2", "label": "MCD", "top": 54, "left": 64},
        {"code": "CAM", "label": "MC", "top": 36, "left": 50},
        {"code": "LW", "label": "EI", "top": 14, "left": 18},
        {"code": "ST", "label": "DC", "top": 8, "left": 50},
        {"code": "RW", "label": "ED", "top": 14, "left": 82},
    ],
    "4-4-2": [
        {"code": "GK", "label": "POR", "top": 90, "left": 50},
        {"code": "LB", "label": "LI", "top": 72, "left": 16},
        {"code": "CB1", "label": "DFC", "top": 76, "left": 36},
        {"code": "CB2", "label": "DFC", "top": 76, "left": 64},
        {"code": "RB", "label": "LD", "top": 72, "left": 84},
        {"code": "LM", "label": "MI", "top": 46, "left": 16},
        {"code": "CM1", "label": "MC", "top": 48, "left": 38},
        {"code": "CM2", "label": "MC", "top": 48, "left": 62},
        {"code": "RM", "label": "MD", "top": 46, "left": 84},
        {"code": "ST1", "label": "DC", "top": 12, "left": 38},
        {"code": "ST2", "label": "DC", "top": 12, "left": 62},
    ],
    "3-5-2": [
        {"code": "GK", "label": "POR", "top": 90, "left": 50},
        {"code": "CB1", "label": "DFC", "top": 76, "left": 30},
        {"code": "CB2", "label": "DFC", "top": 80, "left": 50},
        {"code": "CB3", "label": "DFC", "top": 76, "left": 70},
        {"code": "LM", "label": "MI", "top": 48, "left": 10},
        {"code": "CM1", "label": "MC", "top": 50, "left": 36},
        {"code": "CM2", "label": "MC", "top": 54, "left": 50},
        {"code": "CM3", "label": "MC", "top": 50, "left": 64},
        {"code": "RM", "label": "MD", "top": 48, "left": 90},
        {"code": "ST1", "label": "DC", "top": 12, "left": 38},
        {"code": "ST2", "label": "DC", "top": 12, "left": 62},
    ],
}


def _formacion_valida(nombre):
    return nombre if nombre in FORMACIONES else "4-3-3"


def _obtener_alineacion(db, partido_id, equipo_nombre):
    fila = db.execute(
        "SELECT * FROM alineaciones WHERE partido_id = ? AND equipo = ?",
        (partido_id, equipo_nombre),
    ).fetchone()
    if not fila:
        return None
    jugadores = db.execute(
        """SELECT aj.posicion, aj.titular, j.id, j.nombres, j.apellidos, j.numero_camiseta
           FROM alineacion_jugadores aj JOIN jugadores j ON j.id = aj.jugador_id
           WHERE aj.alineacion_id = ? ORDER BY j.apellidos, j.nombres""",
        (fila["id"],),
    ).fetchall()
    titulares = {j["posicion"]: j for j in jugadores if j["titular"]}
    suplentes = [j for j in jugadores if not j["titular"]]
    return {"formacion": fila["formacion"], "titulares": titulares, "suplentes": suplentes}


def _guardar_alineacion(db, partido_id, equipo_nombre, formacion, titulares, suplentes):
    """titulares: lista de (posicion, jugador_id). suplentes: lista de jugador_id."""
    existente = db.execute(
        "SELECT id FROM alineaciones WHERE partido_id = ? AND equipo = ?",
        (partido_id, equipo_nombre),
    ).fetchone()
    ahora = datetime.now().strftime("%Y-%m-%d %H:%M")
    if existente:
        alineacion_id = existente["id"]
        db.execute("DELETE FROM alineacion_jugadores WHERE alineacion_id = ?", (alineacion_id,))
        db.execute(
            "UPDATE alineaciones SET formacion = ?, fecha_registro = ? WHERE id = ?",
            (formacion, ahora, alineacion_id),
        )
    else:
        returning = " RETURNING id" if USANDO_POSTGRES else ""
        cur = db.execute(
            f"INSERT INTO alineaciones (partido_id, equipo, formacion, fecha_registro) VALUES (?, ?, ?, ?){returning}",
            (partido_id, equipo_nombre, formacion, ahora),
        )
        alineacion_id = cur.fetchone()["id"] if USANDO_POSTGRES else cur.lastrowid

    for posicion, jugador_id in titulares:
        db.execute(
            "INSERT INTO alineacion_jugadores (alineacion_id, jugador_id, posicion, titular) VALUES (?, ?, ?, 1)",
            (alineacion_id, jugador_id, posicion),
        )
    for jugador_id in suplentes:
        db.execute(
            "INSERT INTO alineacion_jugadores (alineacion_id, jugador_id, posicion, titular) VALUES (?, ?, 'SUP', 0)",
            (alineacion_id, jugador_id),
        )
    db.commit()


def _rendimiento_titulares(db, equipo_nombre):
    """Para cada jugador que ha sido titular en alguna alineación guardada de
    este equipo, calcula el récord del equipo (V/E/D) en los partidos ya
    jugados donde ese jugador arrancó de titular."""
    filas = db.execute(
        """SELECT j.id AS jugador_id, j.nombres, j.apellidos,
                  p.equipo_local, p.equipo_visitante, p.goles_local, p.goles_visitante
           FROM alineacion_jugadores aj
           JOIN alineaciones a ON a.id = aj.alineacion_id
           JOIN jugadores j ON j.id = aj.jugador_id
           JOIN partidos p ON p.id = a.partido_id
           WHERE a.equipo = ? AND aj.titular = 1 AND p.jugado = 1
             AND p.goles_local IS NOT NULL AND p.goles_visitante IS NOT NULL""",
        (equipo_nombre,),
    ).fetchall()
    resumen = {}
    for f in filas:
        clave = f["jugador_id"]
        if clave not in resumen:
            resumen[clave] = {"jugador": f"{f['apellidos']} {f['nombres']}", "pj": 0, "v": 0, "e": 0, "d": 0}
        es_local = f["equipo_local"] == equipo_nombre
        gf = f["goles_local"] if es_local else f["goles_visitante"]
        gc = f["goles_visitante"] if es_local else f["goles_local"]
        resumen[clave]["pj"] += 1
        if gf > gc:
            resumen[clave]["v"] += 1
        elif gf < gc:
            resumen[clave]["d"] += 1
        else:
            resumen[clave]["e"] += 1
    filas_resumen = list(resumen.values())
    for r in filas_resumen:
        r["pct"] = round(100 * r["v"] / r["pj"]) if r["pj"] else 0
    filas_resumen.sort(key=lambda r: (-r["pct"], -r["pj"]))
    return filas_resumen


def _conectar():
    if USANDO_POSTGRES:
        return psycopg2.connect(DATABASE_URL, cursor_factory=psycopg2.extras.RealDictCursor)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def get_db():
    if "db" not in g:
        g.db = DBWrapper(_conectar())
    return g.db


@app.teardown_appcontext
def close_db(exception=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def _columnas_existentes(db, tabla):
    if USANDO_POSTGRES:
        filas = db.execute(
            "SELECT column_name FROM information_schema.columns WHERE table_name = ?", (tabla,)
        ).fetchall()
        return [f["column_name"] for f in filas]
    filas = db.execute(f"PRAGMA table_info({tabla})").fetchall()
    return [f[1] for f in filas]


def init_db():
    db = DBWrapper(_conectar())

    tipo_id = "SERIAL PRIMARY KEY" if USANDO_POSTGRES else "INTEGER PRIMARY KEY AUTOINCREMENT"

    db.execute(f"""
        CREATE TABLE IF NOT EXISTS jugadores (
            id {tipo_id},
            cedula TEXT NOT NULL UNIQUE,
            nombres TEXT NOT NULL,
            apellidos TEXT NOT NULL,
            fecha_nacimiento TEXT,
            equipo TEXT NOT NULL,
            categoria TEXT NOT NULL,
            subcategoria TEXT NOT NULL DEFAULT 'Sub 45',
            numero_camiseta TEXT,
            foto TEXT,
            foto_confirmada INTEGER NOT NULL DEFAULT 0,
            cedula_frontal TEXT,
            cedula_reverso TEXT,
            calificado INTEGER NOT NULL DEFAULT 0,
            fecha_registro TEXT NOT NULL
        )
    """)
    jcols = _columnas_existentes(db, "jugadores")
    if "subcategoria" not in jcols:
        db.execute("ALTER TABLE jugadores ADD COLUMN subcategoria TEXT NOT NULL DEFAULT 'Sub 45'")
    if "numero_camiseta" not in jcols:
        db.execute("ALTER TABLE jugadores ADD COLUMN numero_camiseta TEXT")
    if "foto" not in jcols:
        db.execute("ALTER TABLE jugadores ADD COLUMN foto TEXT")
    if "foto_confirmada" not in jcols:
        db.execute("ALTER TABLE jugadores ADD COLUMN foto_confirmada INTEGER NOT NULL DEFAULT 0")
    if "cedula_frontal" not in jcols:
        db.execute("ALTER TABLE jugadores ADD COLUMN cedula_frontal TEXT")
    if "cedula_reverso" not in jcols:
        db.execute("ALTER TABLE jugadores ADD COLUMN cedula_reverso TEXT")
    if "calificado" not in jcols:
        db.execute("ALTER TABLE jugadores ADD COLUMN calificado INTEGER NOT NULL DEFAULT 0")
    if "foto_token" not in jcols:
        db.execute("ALTER TABLE jugadores ADD COLUMN foto_token TEXT")
        for row in db.execute("SELECT id FROM jugadores").fetchall():
            db.execute("UPDATE jugadores SET foto_token = ? WHERE id = ?", (uuid.uuid4().hex, row["id"]))
    if "carnet_impreso" not in jcols:
        db.execute("ALTER TABLE jugadores ADD COLUMN carnet_impreso INTEGER NOT NULL DEFAULT 0")
    if "carnet_valor" not in jcols:
        db.execute("ALTER TABLE jugadores ADD COLUMN carnet_valor REAL")
    if "carnet_fecha" not in jcols:
        db.execute("ALTER TABLE jugadores ADD COLUMN carnet_fecha TEXT")
    if "documentos_fecha" not in jcols:
        db.execute("ALTER TABLE jugadores ADD COLUMN documentos_fecha TEXT")
    if "division" not in jcols:
        db.execute("ALTER TABLE jugadores ADD COLUMN division TEXT NOT NULL DEFAULT ''")

    db.execute(f"""
        CREATE TABLE IF NOT EXISTS equipos (
            id {tipo_id},
            nombre TEXT NOT NULL UNIQUE,
            valor_inscripcion REAL NOT NULL DEFAULT 0,
            abono REAL NOT NULL DEFAULT 0,
            forma_pago TEXT,
            comprobante_pago TEXT,
            usuario TEXT UNIQUE,
            clave TEXT
        )
    """)
    cols = _columnas_existentes(db, "equipos")
    if "valor_inscripcion" not in cols:
        db.execute("ALTER TABLE equipos ADD COLUMN valor_inscripcion REAL NOT NULL DEFAULT 0")
    if "abono" not in cols:
        db.execute("ALTER TABLE equipos ADD COLUMN abono REAL NOT NULL DEFAULT 0")
    if "forma_pago" not in cols:
        db.execute("ALTER TABLE equipos ADD COLUMN forma_pago TEXT")
    if "comprobante_pago" not in cols:
        db.execute("ALTER TABLE equipos ADD COLUMN comprobante_pago TEXT")
    if "usuario" not in cols:
        db.execute("ALTER TABLE equipos ADD COLUMN usuario TEXT")
    if "clave" not in cols:
        db.execute("ALTER TABLE equipos ADD COLUMN clave TEXT")
    if "categoria" not in cols:
        db.execute("ALTER TABLE equipos ADD COLUMN categoria TEXT NOT NULL DEFAULT 'Sub 45'")

    count_row = db.execute("SELECT COUNT(*) c FROM equipos").fetchone()
    count = count_row["c"]
    if count == 0:
        for nombre in EQUIPOS_INICIALES:
            db.execute("INSERT INTO equipos (nombre) VALUES (?)", (nombre,))

    db.execute(f"""
        CREATE TABLE IF NOT EXISTS jornadas (
            id {tipo_id},
            numero INTEGER NOT NULL,
            fecha TEXT,
            categoria TEXT NOT NULL DEFAULT 'Sub 45'
        )
    """)
    if "categoria" not in _columnas_existentes(db, "jornadas"):
        db.execute("ALTER TABLE jornadas ADD COLUMN categoria TEXT NOT NULL DEFAULT 'Sub 45'")

    db.execute(f"""
        CREATE TABLE IF NOT EXISTS partidos (
            id {tipo_id},
            jornada_id INTEGER NOT NULL,
            equipo_local TEXT NOT NULL,
            equipo_visitante TEXT NOT NULL,
            hora TEXT,
            fecha TEXT
        )
    """)
    if "fecha" not in _columnas_existentes(db, "partidos"):
        db.execute("ALTER TABLE partidos ADD COLUMN fecha TEXT")

    db.execute(f"""
        CREATE TABLE IF NOT EXISTS descansos (
            id {tipo_id},
            jornada_id INTEGER NOT NULL,
            equipo TEXT NOT NULL
        )
    """)

    db.execute(f"""
        CREATE TABLE IF NOT EXISTS tarjetas (
            id {tipo_id},
            jugador_id INTEGER NOT NULL,
            jornada_id INTEGER,
            tipo TEXT NOT NULL,
            valor_multa REAL,
            observacion TEXT,
            fecha TEXT NOT NULL
        )
    """)
    if "partido_id" not in _columnas_existentes(db, "tarjetas"):
        db.execute("ALTER TABLE tarjetas ADD COLUMN partido_id INTEGER")
    db.execute(f"""
        CREATE TABLE IF NOT EXISTS sanciones (
            id {tipo_id},
            jugador_id INTEGER NOT NULL,
            motivo TEXT,
            jornadas_sancionado INTEGER NOT NULL DEFAULT 1,
            jornada_desde_id INTEGER,
            valor_multa REAL,
            pagada INTEGER NOT NULL DEFAULT 0,
            fecha TEXT NOT NULL
        )
    """)

    pcols = _columnas_existentes(db, "partidos")
    if "goles_local" not in pcols:
        db.execute("ALTER TABLE partidos ADD COLUMN goles_local INTEGER")
    if "goles_visitante" not in pcols:
        db.execute("ALTER TABLE partidos ADD COLUMN goles_visitante INTEGER")
    if "jugado" not in pcols:
        db.execute("ALTER TABLE partidos ADD COLUMN jugado INTEGER NOT NULL DEFAULT 0")
    if "en_vivo" not in pcols:
        db.execute("ALTER TABLE partidos ADD COLUMN en_vivo INTEGER NOT NULL DEFAULT 0")
    if "arbitro" not in pcols:
        db.execute("ALTER TABLE partidos ADD COLUMN arbitro TEXT")

    db.execute(f"""
        CREATE TABLE IF NOT EXISTS goles (
            id {tipo_id},
            jugador_id INTEGER NOT NULL,
            partido_id INTEGER,
            cantidad INTEGER NOT NULL DEFAULT 1,
            fecha TEXT NOT NULL
        )
    """)

    db.execute(f"""
        CREATE TABLE IF NOT EXISTS categorias (
            id {tipo_id},
            nombre TEXT NOT NULL UNIQUE,
            edad_minima INTEGER
        )
    """)
    if db.execute("SELECT COUNT(*) c FROM categorias").fetchone()["c"] == 0:
        db.execute("INSERT INTO categorias (nombre, edad_minima) VALUES (?, ?)", ("Sub 45", 45))
        db.execute("INSERT INTO categorias (nombre, edad_minima) VALUES (?, ?)", ("Senior", None))

    db.execute(f"""
        CREATE TABLE IF NOT EXISTS divisiones (
            id {tipo_id},
            categoria TEXT NOT NULL,
            nombre TEXT NOT NULL
        )
    """)

    ecols = _columnas_existentes(db, "equipos")
    if "division" not in ecols:
        db.execute("ALTER TABLE equipos ADD COLUMN division TEXT NOT NULL DEFAULT ''")
    jocols = _columnas_existentes(db, "jornadas")
    if "division" not in jocols:
        db.execute("ALTER TABLE jornadas ADD COLUMN division TEXT NOT NULL DEFAULT ''")

    db.execute(f"""
        CREATE TABLE IF NOT EXISTS alineaciones (
            id {tipo_id},
            partido_id INTEGER NOT NULL,
            equipo TEXT NOT NULL,
            formacion TEXT NOT NULL,
            fecha_registro TEXT
        )
    """)
    db.execute(f"""
        CREATE TABLE IF NOT EXISTS alineacion_jugadores (
            id {tipo_id},
            alineacion_id INTEGER NOT NULL,
            jugador_id INTEGER NOT NULL,
            posicion TEXT NOT NULL,
            titular INTEGER NOT NULL DEFAULT 1
        )
    """)

    db.execute(f"""
        CREATE TABLE IF NOT EXISTS vocalia_asistencia (
            id {tipo_id},
            partido_id INTEGER NOT NULL,
            jugador_id INTEGER NOT NULL,
            participo INTEGER NOT NULL DEFAULT 0
        )
    """)
    db.execute(f"""
        CREATE TABLE IF NOT EXISTS vocalia_incidentes (
            id {tipo_id},
            partido_id INTEGER NOT NULL,
            descripcion TEXT NOT NULL,
            fecha TEXT NOT NULL
        )
    """)
    db.execute(f"""
        CREATE TABLE IF NOT EXISTS vocalia_cambios (
            id {tipo_id},
            partido_id INTEGER NOT NULL,
            equipo TEXT NOT NULL,
            jugador_sale_id INTEGER NOT NULL,
            jugador_entra_id INTEGER NOT NULL,
            minuto TEXT,
            fecha TEXT NOT NULL
        )
    """)

    db.commit()
    db.close()


def _to_float(value):
    try:
        return round(float(str(value).replace(",", ".").strip()), 2)
    except (TypeError, ValueError):
        return 0.0


def _contar_juveniles(db, equipo):
    return db.execute(
        "SELECT COUNT(*) c FROM jugadores WHERE equipo = ? AND categoria = ? AND subcategoria = 'Juvenil'",
        (equipo, CATEGORIA_ACTIVA),
    ).fetchone()["c"]


def login_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not session.get("logged_in"):
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return wrapper


def admin_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not session.get("logged_in"):
            return redirect(url_for("login"))
        if session.get("rol") != "admin":
            flash("Esta acción requiere acceso de administrador.")
            return redirect(url_for("index"))
        return f(*args, **kwargs)
    return wrapper


def imprenta_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not session.get("logged_in"):
            return redirect(url_for("login"))
        if session.get("rol") not in ("admin", "imprenta"):
            flash("Esta acción requiere acceso al módulo de carnets.")
            return redirect(url_for("index"))
        return f(*args, **kwargs)
    return wrapper


def calificacion_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not session.get("logged_in"):
            return redirect(url_for("login"))
        if session.get("rol") not in ("admin", "calificacion"):
            flash("Esta acción requiere acceso al módulo de la Comisión de Calificación.")
            return redirect(url_for("index"))
        return f(*args, **kwargs)
    return wrapper


def sanciones_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not session.get("logged_in"):
            return redirect(url_for("login"))
        if session.get("rol") not in ("admin", "sanciones"):
            flash("Esta acción requiere acceso al módulo de Penas y Sanciones.")
            return redirect(url_for("index"))
        return f(*args, **kwargs)
    return wrapper


def tecnica_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not session.get("logged_in"):
            return redirect(url_for("login"))
        if session.get("rol") not in ("admin", "tecnica"):
            flash("Esta acción requiere acceso al módulo de la Comisión Técnica.")
            return redirect(url_for("index"))
        return f(*args, **kwargs)
    return wrapper


def vocalia_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not session.get("logged_in"):
            return redirect(url_for("login"))
        if session.get("rol") not in ("admin", "vocalia"):
            flash("Esta acción requiere acceso al módulo de Vocalía.")
            return redirect(url_for("index"))
        return f(*args, **kwargs)
    return wrapper


def tecnica_o_vocalia_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not session.get("logged_in"):
            return redirect(url_for("login"))
        if session.get("rol") not in ("admin", "tecnica", "vocalia"):
            flash("Esta acción requiere acceso a Comisión Técnica o Vocalía.")
            return redirect(url_for("index"))
        return f(*args, **kwargs)
    return wrapper


def equipo_permitido(equipo_id):
    if session.get("rol") == "admin":
        return True
    return session.get("equipo_id") == equipo_id


def jugador_permitido(db, jugador):
    if session.get("rol") == "admin":
        return True
    fila = db.execute("SELECT id FROM equipos WHERE nombre = ?", (jugador["equipo"],)).fetchone()
    return fila is not None and session.get("equipo_id") == fila["id"]


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        user = request.form.get("usuario", "").strip()
        pw = request.form.get("clave", "").strip()

        if user == ADMIN_USER and pw == ADMIN_PASS:
            session.clear()
            session["logged_in"] = True
            session["rol"] = "admin"
            return redirect(url_for("inscripcion"))

        if IMPRENTA_PASS and user == IMPRENTA_USER and pw == IMPRENTA_PASS:
            session.clear()
            session["logged_in"] = True
            session["rol"] = "imprenta"
            return redirect(url_for("carnets_modulo"))

        if CALIFICACION_PASS and user == CALIFICACION_USER and pw == CALIFICACION_PASS:
            session.clear()
            session["logged_in"] = True
            session["rol"] = "calificacion"
            return redirect(url_for("comision_modulo"))

        if SANCIONES_PASS and user == SANCIONES_USER and pw == SANCIONES_PASS:
            session.clear()
            session["logged_in"] = True
            session["rol"] = "sanciones"
            return redirect(url_for("sanciones_modulo"))

        if TECNICA_PASS and user == TECNICA_USER and pw == TECNICA_PASS:
            session.clear()
            session["logged_in"] = True
            session["rol"] = "tecnica"
            return redirect(url_for("tecnica_modulo"))

        if VOCALIA_PASS and user == VOCALIA_USER and pw == VOCALIA_PASS:
            session.clear()
            session["logged_in"] = True
            session["rol"] = "vocalia"
            return redirect(url_for("vocalia_modulo"))

        db = get_db()
        equipo = db.execute(
            "SELECT * FROM equipos WHERE usuario = ? AND clave = ?", (user, pw)
        ).fetchone()
        if equipo:
            session.clear()
            session["logged_in"] = True
            session["rol"] = "equipo"
            session["equipo_id"] = equipo["id"]
            session["equipo_nombre"] = equipo["nombre"]
            return redirect(url_for("detalle_equipo", equipo_id=equipo["id"]))

        flash("Usuario o clave incorrectos")
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/")
def index():
    if not session.get("logged_in"):
        return redirect(url_for("login"))
    if session.get("rol") == "equipo":
        db = get_db()
        equipo = db.execute("SELECT id FROM equipos WHERE id = ?", (session.get("equipo_id"),)).fetchone()
        if not equipo:
            session.clear()
            flash("Tu sesión ya no es válida. Ingresa nuevamente.")
            return redirect(url_for("login"))
        return redirect(url_for("detalle_equipo", equipo_id=equipo["id"]))
    if session.get("rol") == "imprenta":
        return redirect(url_for("carnets_modulo"))
    if session.get("rol") == "calificacion":
        return redirect(url_for("comision_modulo"))
    if session.get("rol") == "sanciones":
        return redirect(url_for("sanciones_modulo"))
    if session.get("rol") == "tecnica":
        return redirect(url_for("tecnica_modulo"))
    if session.get("rol") == "vocalia":
        return redirect(url_for("vocalia_modulo"))
    return redirect(url_for("inscripcion"))


@app.route("/equipos/agregar", methods=["POST"])
@admin_required
def agregar_equipo():
    db = get_db()
    nombre = request.form.get("nombre", "").strip()
    categoria, division = _categoria_y_division(
        db, request.form.get("categoria", CATEGORIA_ACTIVA), request.form.get("division", "")
    )
    if nombre:
        try:
            db.execute(
                "INSERT INTO equipos (nombre, categoria, division) VALUES (?, ?, ?)",
                (nombre, categoria, division),
            )
            db.commit()
            flash(f"Equipo '{nombre}' agregado.", "ok")
        except IntegrityError:
            db.rollback()
            flash(f"Ya existe un equipo llamado '{nombre}'.")
    return redirect(url_for("inscripcion", categoria=categoria, division=division))


@app.route("/equipos/<int:equipo_id>/renombrar", methods=["POST"])
@admin_required
def renombrar_equipo(equipo_id):
    db = get_db()
    nuevo_nombre = request.form.get("nombre", "").strip()
    if nuevo_nombre:
        actual = db.execute("SELECT nombre FROM equipos WHERE id = ?", (equipo_id,)).fetchone()
        if actual:
            try:
                db.execute("UPDATE equipos SET nombre = ? WHERE id = ?", (nuevo_nombre, equipo_id))
                db.execute("UPDATE jugadores SET equipo = ? WHERE equipo = ?", (nuevo_nombre, actual["nombre"]))
                db.commit()
                flash(f"Equipo renombrado a '{nuevo_nombre}'.", "ok")
            except IntegrityError:
                db.rollback()
                flash(f"Ya existe un equipo llamado '{nuevo_nombre}'.")
    return redirect(url_for("inscripcion"))


@app.route("/equipos/<int:equipo_id>/eliminar", methods=["POST"])
@admin_required
def eliminar_equipo(equipo_id):
    db = get_db()
    equipo = db.execute("SELECT nombre FROM equipos WHERE id = ?", (equipo_id,)).fetchone()
    if equipo:
        jugadores_count = db.execute(
            "SELECT COUNT(*) c FROM jugadores WHERE equipo = ?", (equipo["nombre"],)
        ).fetchone()["c"]
        if jugadores_count > 0:
            flash(f"No se puede eliminar '{equipo['nombre']}': tiene {jugadores_count} jugador(es) inscritos.")
        else:
            db.execute("DELETE FROM equipos WHERE id = ?", (equipo_id,))
            db.commit()
            flash(f"Equipo '{equipo['nombre']}' eliminado.", "ok")
    return redirect(url_for("inscripcion"))


@app.route("/equipos/<int:equipo_id>/credenciales", methods=["POST"])
@admin_required
def credenciales_equipo(equipo_id):
    db = get_db()
    usuario = request.form.get("usuario", "").strip().lower()
    clave = request.form.get("clave", "").strip()
    if not (usuario and clave):
        flash("Usuario y clave son obligatorios.")
    else:
        try:
            db.execute("UPDATE equipos SET usuario = ?, clave = ? WHERE id = ?", (usuario, clave, equipo_id))
            db.commit()
            flash(f"Acceso del equipo actualizado: usuario '{usuario}'.", "ok")
        except IntegrityError:
            db.rollback()
            flash(f"El usuario '{usuario}' ya está en uso por otro equipo.")
    return redirect(url_for("detalle_equipo", equipo_id=equipo_id))


@app.route("/mi_equipo/cambiar_clave", methods=["POST"])
@login_required
def cambiar_clave_equipo():
    """El delegado del equipo cambia su propia clave de acceso, a la que
    prefiera. El administrador conserva acceso total de todas formas (puede
    ver y volver a cambiar la clave de cualquier equipo desde su panel)."""
    if session.get("rol") != "equipo":
        flash("Esta opción es solo para el acceso de delegado de equipo.")
        return redirect(url_for("index"))

    equipo_id = session.get("equipo_id")
    nueva_clave = request.form.get("nueva_clave", "").strip()
    confirmar_clave = request.form.get("confirmar_clave", "").strip()

    if not nueva_clave or len(nueva_clave) < 4:
        flash("La nueva clave debe tener al menos 4 caracteres.")
    elif nueva_clave != confirmar_clave:
        flash("Las dos claves no coinciden.")
    else:
        db = get_db()
        db.execute("UPDATE equipos SET clave = ? WHERE id = ?", (nueva_clave, equipo_id))
        db.commit()
        flash("Tu clave de acceso se actualizó correctamente.", "ok")

    return redirect(url_for("detalle_equipo", equipo_id=equipo_id))


@app.route("/equipos/<int:equipo_id>/credenciales/quitar", methods=["POST"])
@admin_required
def quitar_credenciales_equipo(equipo_id):
    db = get_db()
    db.execute("UPDATE equipos SET usuario = NULL, clave = NULL WHERE id = ?", (equipo_id,))
    db.commit()
    flash("Se quitó el acceso del delegado para este equipo.", "ok")
    return redirect(url_for("detalle_equipo", equipo_id=equipo_id))


@app.route("/equipo/<int:equipo_id>", methods=["GET"])
@login_required
def detalle_equipo(equipo_id):
    db = get_db()
    equipo = db.execute("SELECT * FROM equipos WHERE id = ?", (equipo_id,)).fetchone()
    if not equipo:
        flash("Equipo no encontrado.")
        return redirect(url_for("index"))

    if not equipo_permitido(equipo_id):
        flash("No tienes acceso a ese equipo.")
        return redirect(url_for("index"))

    jugadores = db.execute(
        "SELECT * FROM jugadores WHERE equipo = ? AND categoria = ? ORDER BY apellidos",
        (equipo["nombre"], equipo["categoria"]),
    ).fetchall()

    saldo = equipo["valor_inscripcion"] - equipo["abono"]
    juveniles_count = _contar_juveniles(db, equipo["nombre"])

    partidos_equipo = []
    for jr in db.execute(
        "SELECT * FROM jornadas WHERE categoria = ? AND division = ? ORDER BY numero, id",
        (equipo["categoria"], equipo["division"]),
    ).fetchall():
        partido = db.execute(
            "SELECT * FROM partidos WHERE jornada_id = ? AND (equipo_local = ? OR equipo_visitante = ?)",
            (jr["id"], equipo["nombre"], equipo["nombre"]),
        ).fetchone()
        descansa = db.execute(
            "SELECT * FROM descansos WHERE jornada_id = ? AND equipo = ?", (jr["id"], equipo["nombre"])
        ).fetchone()
        if partido:
            rival = partido["equipo_visitante"] if partido["equipo_local"] == equipo["nombre"] else partido["equipo_local"]
            local = partido["equipo_local"] == equipo["nombre"]
            partidos_equipo.append({
                "jornada": jr, "rival": rival, "local": local,
                "hora": partido["hora"], "fecha": partido["fecha"],
                "partido_id": partido["id"],
            })
        elif descansa:
            partidos_equipo.append({"jornada": jr, "rival": None, "local": None, "hora": None, "fecha": None, "partido_id": None})

    return render_template(
        "detalle_equipo.html",
        equipo=equipo,
        jugadores=jugadores,
        saldo=saldo,
        cupo_maximo=CUPO_MAXIMO_EQUIPO,
        categoria=equipo["categoria"],
        juveniles_count=juveniles_count,
        cupo_maximo_juvenil=CUPO_MAXIMO_JUVENIL,
        partidos_equipo=partidos_equipo,
        rendimiento_titulares=_rendimiento_titulares(db, equipo["nombre"]),
    )


@app.route("/equipo/<int:equipo_id>/partido/<int:partido_id>/alineacion", methods=["GET"])
@login_required
def alineacion_equipo(equipo_id, partido_id):
    if not equipo_permitido(equipo_id):
        flash("No tienes acceso a ese equipo.")
        return redirect(url_for("index"))
    db = get_db()
    equipo = db.execute("SELECT * FROM equipos WHERE id = ?", (equipo_id,)).fetchone()
    if not equipo:
        flash("Equipo no encontrado.")
        return redirect(url_for("index"))

    partido = db.execute("SELECT * FROM partidos WHERE id = ?", (partido_id,)).fetchone()
    if not partido or equipo["nombre"] not in (partido["equipo_local"], partido["equipo_visitante"]):
        flash("Ese partido no corresponde a este equipo.")
        return redirect(url_for("detalle_equipo", equipo_id=equipo_id))
    rival = partido["equipo_visitante"] if partido["equipo_local"] == equipo["nombre"] else partido["equipo_local"]

    existente = _obtener_alineacion(db, partido_id, equipo["nombre"])
    formacion = _formacion_valida(request.args.get("formacion", existente["formacion"] if existente else "4-3-3"))
    slots = FORMACIONES[formacion]

    jugadores = db.execute(
        "SELECT * FROM jugadores WHERE equipo = ? ORDER BY apellidos, nombres", (equipo["nombre"],)
    ).fetchall()

    seleccion = {}
    suplentes_ids = set()
    if existente:
        suplentes_ids = {j["id"] for j in existente["suplentes"]}
        if existente["formacion"] == formacion:
            seleccion = {pos: j["id"] for pos, j in existente["titulares"].items()}
        else:
            # cambió de formación: los que eran titulares pasan a quedar disponibles otra vez
            suplentes_ids |= {j["id"] for j in existente["titulares"].values()}

    return render_template(
        "alineacion_equipo.html",
        equipo=equipo, partido=partido, rival=rival,
        formaciones=list(FORMACIONES.keys()), formacion=formacion, slots=slots,
        jugadores=jugadores, seleccion=seleccion, suplentes_ids=suplentes_ids,
        tiene_alineacion=existente is not None,
    )


@app.route("/equipo/<int:equipo_id>/partido/<int:partido_id>/alineacion/guardar", methods=["POST"])
@login_required
def guardar_alineacion(equipo_id, partido_id):
    if not equipo_permitido(equipo_id):
        flash("No tienes acceso a ese equipo.")
        return redirect(url_for("index"))
    db = get_db()
    equipo = db.execute("SELECT * FROM equipos WHERE id = ?", (equipo_id,)).fetchone()
    partido = db.execute("SELECT * FROM partidos WHERE id = ?", (partido_id,)).fetchone()
    if not equipo or not partido or equipo["nombre"] not in (partido["equipo_local"], partido["equipo_visitante"]):
        flash("Ese partido no corresponde a este equipo.")
        return redirect(url_for("index"))

    formacion = _formacion_valida(request.form.get("formacion", "4-3-3"))
    slots = FORMACIONES[formacion]

    ids_validos = {
        j["id"] for j in db.execute("SELECT id FROM jugadores WHERE equipo = ?", (equipo["nombre"],)).fetchall()
    }

    titulares = []
    usados = set()
    for slot in slots:
        raw = request.form.get(f"pos_{slot['code']}", "").strip()
        if not raw.isdigit():
            continue
        jugador_id = int(raw)
        if jugador_id not in ids_validos or jugador_id in usados:
            continue
        titulares.append((slot["code"], jugador_id))
        usados.add(jugador_id)

    suplentes = []
    for raw in request.form.getlist("suplentes"):
        if raw.isdigit():
            jugador_id = int(raw)
            if jugador_id in ids_validos and jugador_id not in usados:
                suplentes.append(jugador_id)
                usados.add(jugador_id)

    if not titulares:
        flash("Selecciona al menos un jugador titular.")
        return redirect(url_for("alineacion_equipo", equipo_id=equipo_id, partido_id=partido_id, formacion=formacion))

    _guardar_alineacion(db, partido_id, equipo["nombre"], formacion, titulares, suplentes)
    flash("Alineación guardada.", "ok")
    return redirect(url_for("alineacion_equipo", equipo_id=equipo_id, partido_id=partido_id, formacion=formacion))


@app.route("/publico/partido/<int:partido_id>/alineaciones")
def publico_alineaciones(partido_id):
    db = get_db()
    partido = db.execute("SELECT * FROM partidos WHERE id = ?", (partido_id,)).fetchone()
    if not partido:
        flash("Partido no encontrado.")
        return redirect(url_for("publico_inicio"))
    jornada = db.execute("SELECT * FROM jornadas WHERE id = ?", (partido["jornada_id"],)).fetchone()

    def _armar(equipo_nombre):
        datos = _obtener_alineacion(db, partido_id, equipo_nombre)
        if not datos:
            return None
        slots = FORMACIONES.get(datos["formacion"], FORMACIONES["4-3-3"])
        return {
            "formacion": datos["formacion"], "slots": slots,
            "titulares": datos["titulares"], "suplentes": datos["suplentes"],
        }

    def _tarjetas(equipo_nombre):
        # Cada equipo juega a lo más un partido por jornada, así que las
        # tarjetas de esa jornada para jugadores de este equipo son,
        # necesariamente, las de este encuentro.
        if not jornada:
            return []
        return db.execute(
            """SELECT t.tipo, j.nombres, j.apellidos, j.numero_camiseta
               FROM tarjetas t JOIN jugadores j ON j.id = t.jugador_id
               WHERE t.jornada_id = ? AND j.equipo = ?
               ORDER BY t.fecha""",
            (jornada["id"], equipo_nombre),
        ).fetchall()

    asistencia_local = db.execute(
        """SELECT j.nombres, j.apellidos, j.numero_camiseta
           FROM vocalia_asistencia va JOIN jugadores j ON j.id = va.jugador_id
           WHERE va.partido_id = ? AND va.participo = 1 AND j.equipo = ?
           ORDER BY j.apellidos, j.nombres""",
        (partido_id, partido["equipo_local"]),
    ).fetchall()
    asistencia_visitante = db.execute(
        """SELECT j.nombres, j.apellidos, j.numero_camiseta
           FROM vocalia_asistencia va JOIN jugadores j ON j.id = va.jugador_id
           WHERE va.partido_id = ? AND va.participo = 1 AND j.equipo = ?
           ORDER BY j.apellidos, j.nombres""",
        (partido_id, partido["equipo_visitante"]),
    ).fetchall()
    incidentes = db.execute(
        "SELECT * FROM vocalia_incidentes WHERE partido_id = ? ORDER BY fecha", (partido_id,)
    ).fetchall()
    cambios = db.execute(
        """SELECT c.*, js.apellidos AS sale_apellidos, js.nombres AS sale_nombres, js.numero_camiseta AS sale_numero,
                  je.apellidos AS entra_apellidos, je.nombres AS entra_nombres, je.numero_camiseta AS entra_numero
           FROM vocalia_cambios c
           JOIN jugadores js ON js.id = c.jugador_sale_id
           JOIN jugadores je ON je.id = c.jugador_entra_id
           WHERE c.partido_id = ? ORDER BY c.fecha""",
        (partido_id,),
    ).fetchall()

    return render_template(
        "publico_alineaciones.html",
        partido=partido, jornada=jornada,
        local=_armar(partido["equipo_local"]), visitante=_armar(partido["equipo_visitante"]),
        tarjetas_local=_tarjetas(partido["equipo_local"]),
        tarjetas_visitante=_tarjetas(partido["equipo_visitante"]),
        asistencia_local=asistencia_local, asistencia_visitante=asistencia_visitante,
        incidentes=incidentes, cambios=cambios,
    )


def _exportar_jugadores_excel(jugadores, nombre_archivo, incluir_equipo=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Jugadores"
    encabezados = ["Cédula", "Nombres", "Apellidos", "Fecha nacimiento", "Campeonato"]
    if incluir_equipo:
        encabezados.insert(0, "Equipo")
    encabezados += ["Número camiseta", "Registrado"]
    ws.append(encabezados)

    for j in jugadores:
        categoria = j["categoria"] + (" - Juvenil" if j["subcategoria"] == "Juvenil" else "")
        fila = [j["cedula"], j["nombres"], j["apellidos"], j["fecha_nacimiento"] or "", categoria]
        if incluir_equipo:
            fila.insert(0, j["equipo"])
        fila += [j["numero_camiseta"] or "", j["fecha_registro"]]
        ws.append(fila)

    for i, _ in enumerate(encabezados, start=1):
        ws.column_dimensions[chr(64 + i)].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=nombre_archivo,
    )


@app.route("/equipo/<int:equipo_id>/exportar")
@login_required
def exportar_equipo(equipo_id):
    db = get_db()
    equipo = db.execute("SELECT * FROM equipos WHERE id = ?", (equipo_id,)).fetchone()
    if not equipo:
        flash("Equipo no encontrado.")
        return redirect(url_for("index"))
    if not equipo_permitido(equipo_id):
        flash("No tienes acceso a ese equipo.")
        return redirect(url_for("index"))

    jugadores = db.execute(
        "SELECT * FROM jugadores WHERE equipo = ? AND categoria = ? ORDER BY apellidos",
        (equipo["nombre"], equipo["categoria"]),
    ).fetchall()
    nombre_archivo = f"jugadores_{equipo['nombre'].replace(' ', '_')}.xlsx"
    return _exportar_jugadores_excel(jugadores, nombre_archivo)


@app.route("/jugadores_liga")
@login_required
def jugadores_liga():
    """Lista de todos los jugadores inscritos en la liga, de todos los
    equipos, visible tanto para el admin como para los delegados -- para
    que cualquiera pueda revisar quién está inscrito y en qué equipo."""
    db = get_db()
    categoria, division = _categoria_y_division(
        db, request.args.get("categoria", CATEGORIA_ACTIVA), request.args.get("division", "")
    )
    jugadores = db.execute(
        "SELECT * FROM jugadores WHERE categoria = ? AND division = ? ORDER BY equipo, apellidos",
        (categoria, division),
    ).fetchall()
    equipos_distintos = sorted(set(j["equipo"] for j in jugadores))
    return render_template(
        "jugadores_liga.html",
        jugadores=jugadores,
        categoria=categoria,
        categorias_liga=_categorias_liga(db),
        division=division,
        divisiones=_divisiones_de_categoria(db, categoria),
        total_equipos=len(equipos_distintos),
    )


@app.route("/exportar_general")
@admin_required
def exportar_general():
    db = get_db()
    categoria, division = _categoria_y_division(
        db, request.args.get("categoria", CATEGORIA_ACTIVA), request.args.get("division", "")
    )
    jugadores = db.execute(
        "SELECT * FROM jugadores WHERE categoria = ? AND division = ? ORDER BY equipo, apellidos",
        (categoria, division),
    ).fetchall()
    sufijo = f"{categoria}_{division}" if division else categoria
    nombre_archivo = f"jugadores_liga_oyambarillo_{sufijo.replace(' ', '_')}.xlsx"
    return _exportar_jugadores_excel(jugadores, nombre_archivo, incluir_equipo=True)


FORMAS_PAGO_VALIDAS = {"Efectivo", "Depósito", "Transferencia"}


@app.route("/equipo/<int:equipo_id>/pago", methods=["POST"])
@admin_required
def actualizar_pago_equipo(equipo_id):
    db = get_db()
    valor_inscripcion = _to_float(request.form.get("valor_inscripcion", "0"))
    abono = _to_float(request.form.get("abono", "0"))
    forma_pago = request.form.get("forma_pago", "Efectivo").strip()
    if forma_pago not in FORMAS_PAGO_VALIDAS:
        forma_pago = "Efectivo"
    comprobante_pago = request.form.get("comprobante_pago", "").strip()
    db.execute(
        "UPDATE equipos SET valor_inscripcion = ?, abono = ?, forma_pago = ?, comprobante_pago = ? WHERE id = ?",
        (valor_inscripcion, abono, forma_pago, comprobante_pago, equipo_id),
    )
    db.commit()
    flash("Datos de pago actualizados.", "ok")
    return redirect(url_for("detalle_equipo", equipo_id=equipo_id))


def _guardar_imagen_subida(archivo):
    if not archivo or not archivo.filename:
        return None
    ext = os.path.splitext(archivo.filename)[1].lower()
    if ext not in (".jpg", ".jpeg", ".png"):
        return None
    nuevo_nombre = f"{uuid.uuid4().hex}.jpg"
    try:
        img = Image.open(archivo)
        img = ImageOps.exif_transpose(img).convert("RGB")
        img.thumbnail((1600, 1600), Image.LANCZOS)
        img.save(os.path.join(FOTOS_DIR, nuevo_nombre), format="JPEG", quality=88)
    except Exception:
        return None
    return nuevo_nombre


def _subcategoria_por_nacimiento(fecha_nacimiento, categoria=CATEGORIA_ACTIVA):
    # "Juvenil" es una excepción histórica exclusiva de Sub 45 (nacidos en 1982).
    # Cualquier otra categoría (Senior o una nueva que arme el admin) no tiene sub-tier.
    if categoria != "Sub 45":
        return categoria
    return "Juvenil" if (fecha_nacimiento or "").strip().startswith("1982") else "Sub 45"


def _jugador_califica(db, fecha_nacimiento, categoria=CATEGORIA_ACTIVA):
    """Sub 45: solo califica quien nace en 1982 (Juvenil, cupo aparte) o cumple
    la edad mínima de la categoría en el año actual de la temporada. Las demás
    categorías usan la edad mínima configurada (o ninguna, si no se definió)."""
    fecha_nacimiento = (fecha_nacimiento or "").strip()
    if categoria == "Sub 45" and fecha_nacimiento.startswith("1982"):
        return True

    edad_minima = _edad_minima_categoria(db, categoria)
    if edad_minima is None:
        return True

    try:
        anio_nacimiento = int(fecha_nacimiento[:4])
    except (ValueError, IndexError):
        return False
    return (date.today().year - anio_nacimiento) >= edad_minima


def _insertar_jugador(db, equipo_nombre, cedula, nombres, apellidos, fecha_nacimiento, categoria,
                       numero_camiseta="", foto=None, cedula_frontal=None, cedula_reverso=None):
    categoria = _categoria_valida(db, categoria)
    subcategoria = _subcategoria_por_nacimiento(fecha_nacimiento, categoria)

    if not (cedula and nombres and apellidos):
        return None, "Cédula, nombres y apellidos son obligatorios."

    if not _jugador_califica(db, fecha_nacimiento, categoria):
        edad_minima = _edad_minima_categoria(db, categoria)
        if categoria == "Sub 45":
            return None, ("Cédula no califica: según la fecha de nacimiento, no cumple 45 años este año "
                           "ni nació en 1982 (Juvenil). No corresponde a ningún campeonato de esta liga.")
        return None, (f"Cédula no califica: el campeonato {categoria} exige {edad_minima} años cumplidos "
                       "este año. No corresponde a este campeonato.")

    count = db.execute(
        "SELECT COUNT(*) c FROM jugadores WHERE equipo = ? AND categoria = ?",
        (equipo_nombre, categoria),
    ).fetchone()["c"]
    if count >= CUPO_MAXIMO_EQUIPO:
        return None, f"El equipo {equipo_nombre} ya alcanzó el cupo máximo de {CUPO_MAXIMO_EQUIPO} jugadores."

    if subcategoria == "Juvenil" and _contar_juveniles(db, equipo_nombre) >= CUPO_MAXIMO_JUVENIL:
        return None, f"El equipo {equipo_nombre} ya alcanzó el cupo máximo de {CUPO_MAXIMO_JUVENIL} jugadores Juvenil."

    ahora = datetime.now().strftime("%Y-%m-%d %H:%M")
    documentos_fecha = ahora if (foto or cedula_frontal or cedula_reverso) else None
    equipo_row = db.execute("SELECT division FROM equipos WHERE nombre = ?", (equipo_nombre,)).fetchone()
    division = equipo_row["division"] if equipo_row else ""

    returning = " RETURNING id" if USANDO_POSTGRES else ""
    try:
        cur = db.execute(
            f"""INSERT INTO jugadores
               (cedula, nombres, apellidos, fecha_nacimiento, equipo, categoria, division, subcategoria, numero_camiseta,
                foto, cedula_frontal, cedula_reverso, foto_token, fecha_registro, documentos_fecha)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?){returning}""",
            (cedula, nombres, apellidos, fecha_nacimiento, equipo_nombre, categoria, division,
             subcategoria, numero_camiseta, foto, cedula_frontal, cedula_reverso,
             uuid.uuid4().hex, ahora, documentos_fecha),
        )
        db.commit()
        if USANDO_POSTGRES:
            return cur.fetchone()["id"], None
        return cur.lastrowid, None
    except IntegrityError:
        db.rollback()
        return None, f"Ya existe un jugador registrado con la cédula {cedula}."


@app.route("/api/cedula_existe")
@login_required
def api_cedula_existe():
    """Consulta rápida (para avisar en el formulario antes de guardar) si
    ya existe un jugador registrado con esa cédula, en cualquier equipo —
    la cédula es única en todo el sistema."""
    cedula = request.args.get("cedula", "").strip()
    if not cedula:
        return jsonify({"existe": False})
    db = get_db()
    jugador = db.execute(
        "SELECT nombres, apellidos, equipo FROM jugadores WHERE cedula = ?", (cedula,)
    ).fetchone()
    if not jugador:
        return jsonify({"existe": False})
    return jsonify({
        "existe": True,
        "nombres": jugador["nombres"],
        "apellidos": jugador["apellidos"],
        "equipo": jugador["equipo"],
    })


@app.route("/equipo/<int:equipo_id>/agregar_jugador", methods=["POST"])
@login_required
def agregar_jugador_equipo(equipo_id):
    db = get_db()
    equipo = db.execute("SELECT * FROM equipos WHERE id = ?", (equipo_id,)).fetchone()
    if not equipo:
        flash("Equipo no encontrado.")
        return redirect(url_for("index"))
    if not equipo_permitido(equipo_id):
        flash("No tienes acceso a ese equipo.")
        return redirect(url_for("index"))

    foto = _guardar_imagen_subida(request.files.get("foto"))
    cedula_frontal = _guardar_imagen_subida(request.files.get("cedula_frontal"))
    cedula_reverso = _guardar_imagen_subida(request.files.get("cedula_reverso"))

    jugador_id, error = _insertar_jugador(
        db,
        equipo["nombre"],
        request.form.get("cedula", "").strip(),
        request.form.get("nombres", "").strip(),
        request.form.get("apellidos", "").strip(),
        request.form.get("fecha_nacimiento", "").strip(),
        equipo["categoria"],
        request.form.get("numero_camiseta", "").strip(),
        foto=foto,
        cedula_frontal=cedula_frontal,
        cedula_reverso=cedula_reverso,
    )
    if error:
        flash(error)
    else:
        flash("Jugador registrado correctamente.", "ok")
        return redirect(url_for("ficha_jugador", jugador_id=jugador_id))

    return redirect(url_for("detalle_equipo", equipo_id=equipo_id))


@app.route("/inscripcion", methods=["GET"])
@admin_required
def inscripcion():
    db = get_db()
    categoria, division = _categoria_y_division(
        db, request.args.get("categoria", CATEGORIA_ACTIVA), request.args.get("division", "")
    )

    equipos_rows = db.execute(
        "SELECT * FROM equipos WHERE categoria = ? AND division = ? ORDER BY nombre", (categoria, division)
    ).fetchall()

    cupos = {}
    for equipo in equipos_rows:
        c = db.execute(
            "SELECT COUNT(*) c FROM jugadores WHERE equipo = ? AND categoria = ?",
            (equipo["nombre"], categoria),
        ).fetchone()["c"]
        cupos[equipo["nombre"]] = c

    return render_template(
        "inscripcion.html",
        equipos=equipos_rows,
        cupos=cupos,
        cupo_maximo=CUPO_MAXIMO_EQUIPO,
        categoria=categoria,
        categorias_liga=_categorias_liga(db),
        division=division,
        divisiones=_divisiones_de_categoria(db, categoria),
        subcategorias=SUBCATEGORIAS,
    )


@app.route("/categorias")
@admin_required
def categorias_modulo():
    db = get_db()
    filas = db.execute("SELECT * FROM categorias ORDER BY id").fetchall()
    categorias = []
    for f in filas:
        equipos_count = db.execute(
            "SELECT COUNT(*) c FROM equipos WHERE categoria = ?", (f["nombre"],)
        ).fetchone()["c"]
        divisiones_rows = db.execute(
            "SELECT * FROM divisiones WHERE categoria = ? ORDER BY id", (f["nombre"],)
        ).fetchall()
        divisiones = []
        for d in divisiones_rows:
            d_equipos_count = db.execute(
                "SELECT COUNT(*) c FROM equipos WHERE categoria = ? AND division = ?",
                (f["nombre"], d["nombre"]),
            ).fetchone()["c"]
            divisiones.append({"fila": d, "equipos_count": d_equipos_count})
        categorias.append({"fila": f, "equipos_count": equipos_count, "divisiones": divisiones})
    return render_template("categorias_modulo.html", categorias=categorias)


@app.route("/categorias/agregar", methods=["POST"])
@admin_required
def agregar_categoria():
    db = get_db()
    nombre = request.form.get("nombre", "").strip()
    edad_minima_raw = request.form.get("edad_minima", "").strip()
    edad_minima = int(edad_minima_raw) if edad_minima_raw.isdigit() else None

    if not nombre:
        flash("Ingresa un nombre para el campeonato.")
        return redirect(url_for("categorias_modulo"))

    try:
        db.execute("INSERT INTO categorias (nombre, edad_minima) VALUES (?, ?)", (nombre, edad_minima))
        db.commit()
        flash(f"Campeonato '{nombre}' agregado.", "ok")
    except IntegrityError:
        db.rollback()
        flash(f"Ya existe un campeonato llamado '{nombre}'.")
    return redirect(url_for("categorias_modulo"))


@app.route("/categorias/<int:categoria_id>/editar", methods=["POST"])
@admin_required
def editar_categoria(categoria_id):
    db = get_db()
    edad_minima_raw = request.form.get("edad_minima", "").strip()
    edad_minima = int(edad_minima_raw) if edad_minima_raw.isdigit() else None
    db.execute("UPDATE categorias SET edad_minima = ? WHERE id = ?", (edad_minima, categoria_id))
    db.commit()
    flash("Campeonato actualizado.", "ok")
    return redirect(url_for("categorias_modulo"))


@app.route("/categorias/<int:categoria_id>/eliminar", methods=["POST"])
@admin_required
def eliminar_categoria(categoria_id):
    db = get_db()
    categoria = db.execute("SELECT * FROM categorias WHERE id = ?", (categoria_id,)).fetchone()
    if not categoria:
        flash("Campeonato no encontrado.")
        return redirect(url_for("categorias_modulo"))

    equipos_count = db.execute(
        "SELECT COUNT(*) c FROM equipos WHERE categoria = ?", (categoria["nombre"],)
    ).fetchone()["c"]
    if equipos_count > 0:
        flash(f"No se puede eliminar '{categoria['nombre']}': todavía tiene {equipos_count} equipo(s). Elimina o cambia esos equipos primero.")
        return redirect(url_for("categorias_modulo"))

    if len(_categorias_liga(db)) <= 1:
        flash("No se puede eliminar el último campeonato de la liga.")
        return redirect(url_for("categorias_modulo"))

    db.execute("DELETE FROM divisiones WHERE categoria = ?", (categoria["nombre"],))
    db.execute("DELETE FROM categorias WHERE id = ?", (categoria_id,))
    db.commit()
    flash(f"Campeonato '{categoria['nombre']}' eliminado.", "ok")
    return redirect(url_for("categorias_modulo"))


@app.route("/categorias/<int:categoria_id>/divisiones/agregar", methods=["POST"])
@admin_required
def agregar_division(categoria_id):
    db = get_db()
    categoria = db.execute("SELECT * FROM categorias WHERE id = ?", (categoria_id,)).fetchone()
    if not categoria:
        flash("Campeonato no encontrado.")
        return redirect(url_for("categorias_modulo"))

    nombre = request.form.get("nombre", "").strip()
    if not nombre:
        flash("Ingresa un nombre para la categoría.")
        return redirect(url_for("categorias_modulo"))

    ya_existe = db.execute(
        "SELECT 1 FROM divisiones WHERE categoria = ? AND nombre = ?", (categoria["nombre"], nombre)
    ).fetchone()
    if ya_existe:
        flash(f"'{categoria['nombre']}' ya tiene una categoría llamada '{nombre}'.")
        return redirect(url_for("categorias_modulo"))

    db.execute("INSERT INTO divisiones (categoria, nombre) VALUES (?, ?)", (categoria["nombre"], nombre))
    db.commit()
    flash(f"Categoría '{nombre}' agregada a {categoria['nombre']}.", "ok")
    return redirect(url_for("categorias_modulo"))


@app.route("/divisiones/<int:division_id>/eliminar", methods=["POST"])
@admin_required
def eliminar_division(division_id):
    db = get_db()
    division = db.execute("SELECT * FROM divisiones WHERE id = ?", (division_id,)).fetchone()
    if not division:
        flash("Categoría no encontrada.")
        return redirect(url_for("categorias_modulo"))

    equipos_count = db.execute(
        "SELECT COUNT(*) c FROM equipos WHERE categoria = ? AND division = ?",
        (division["categoria"], division["nombre"]),
    ).fetchone()["c"]
    if equipos_count > 0:
        flash(f"No se puede eliminar '{division['nombre']}': todavía tiene {equipos_count} equipo(s).")
        return redirect(url_for("categorias_modulo"))

    db.execute("DELETE FROM divisiones WHERE id = ?", (division_id,))
    db.commit()
    flash(f"Categoría '{division['nombre']}' eliminada.", "ok")
    return redirect(url_for("categorias_modulo"))


def _calcular_edad(fecha_nacimiento):
    try:
        y, m, d = [int(p) for p in fecha_nacimiento.split("-")]
        nacimiento = date(y, m, d)
    except (ValueError, AttributeError):
        return None
    hoy = date.today()
    edad = hoy.year - nacimiento.year - ((hoy.month, hoy.day) < (nacimiento.month, nacimiento.day))
    return edad


@app.route("/jugador/<int:jugador_id>", methods=["GET"])
@login_required
def ficha_jugador(jugador_id):
    db = get_db()
    jugador = db.execute("SELECT * FROM jugadores WHERE id = ?", (jugador_id,)).fetchone()
    if not jugador:
        flash("Jugador no encontrado.")
        return redirect(url_for("index"))
    if not jugador_permitido(db, jugador):
        flash("No tienes acceso a ese jugador.")
        return redirect(url_for("index"))
    edad = _calcular_edad(jugador["fecha_nacimiento"])
    equipo_row = db.execute("SELECT id FROM equipos WHERE nombre = ?", (jugador["equipo"],)).fetchone()
    equipo_id = equipo_row["id"] if equipo_row else None
    return render_template("ficha_jugador.html", jugador=jugador, edad=edad, subcategorias=SUBCATEGORIAS, equipo_id=equipo_id)


@app.route("/jugador/<int:jugador_id>/actualizar", methods=["POST"])
@login_required
def actualizar_jugador(jugador_id):
    db = get_db()
    jugador = db.execute("SELECT * FROM jugadores WHERE id = ?", (jugador_id,)).fetchone()
    if not jugador:
        flash("Jugador no encontrado.")
        return redirect(url_for("index"))
    if not jugador_permitido(db, jugador):
        flash("No tienes acceso a ese jugador.")
        return redirect(url_for("index"))

    if jugador["calificado"] and session.get("rol") != "admin":
        flash("Este jugador ya fue calificado y no se puede modificar. Contacta a la Comisión de Calificación.")
        return redirect(url_for("ficha_jugador", jugador_id=jugador_id))

    nombres = request.form.get("nombres", "").strip() or jugador["nombres"]
    apellidos = request.form.get("apellidos", "").strip() or jugador["apellidos"]
    fecha_nacimiento = request.form.get("fecha_nacimiento", "").strip() or jugador["fecha_nacimiento"]
    numero_camiseta = request.form.get("numero_camiseta", "").strip()
    subcategoria = _subcategoria_por_nacimiento(fecha_nacimiento, jugador["categoria"])

    if not _jugador_califica(db, fecha_nacimiento, jugador["categoria"]):
        flash("Cédula no califica: no cumple con la edad mínima de este campeonato.")
        return redirect(url_for("ficha_jugador", jugador_id=jugador_id))

    if (subcategoria == "Juvenil" and jugador["subcategoria"] != "Juvenil"
            and _contar_juveniles(db, jugador["equipo"]) >= CUPO_MAXIMO_JUVENIL):
        flash(f"No se puede cambiar la fecha: el equipo ya alcanzó el cupo máximo de {CUPO_MAXIMO_JUVENIL} jugadores Juvenil.")
        return redirect(url_for("ficha_jugador", jugador_id=jugador_id))

    nuevo_foto = _guardar_imagen_subida(request.files.get("foto"))
    nuevo_cedula_frontal = _guardar_imagen_subida(request.files.get("cedula_frontal"))
    nuevo_cedula_reverso = _guardar_imagen_subida(request.files.get("cedula_reverso"))
    foto_nombre = nuevo_foto or jugador["foto"]
    cedula_frontal = nuevo_cedula_frontal or jugador["cedula_frontal"]
    cedula_reverso = nuevo_cedula_reverso or jugador["cedula_reverso"]
    documentos_fecha = (
        datetime.now().strftime("%Y-%m-%d %H:%M")
        if (nuevo_foto or nuevo_cedula_frontal or nuevo_cedula_reverso)
        else jugador["documentos_fecha"]
    )

    db.execute(
        """UPDATE jugadores SET nombres = ?, apellidos = ?, fecha_nacimiento = ?, subcategoria = ?,
           numero_camiseta = ?, foto = ?, cedula_frontal = ?, cedula_reverso = ?, documentos_fecha = ? WHERE id = ?""",
        (nombres, apellidos, fecha_nacimiento, subcategoria, numero_camiseta, foto_nombre,
         cedula_frontal, cedula_reverso, documentos_fecha, jugador_id),
    )
    db.commit()
    flash("Ficha del jugador actualizada.", "ok")
    return redirect(url_for("ficha_jugador", jugador_id=jugador_id))


CAMPOS_FOTO_VALIDOS = {"foto", "cedula_frontal", "cedula_reverso"}


@app.route("/jugador/<int:jugador_id>/quitar_foto/<campo>", methods=["POST"])
@login_required
def quitar_foto_jugador(jugador_id, campo):
    if campo not in CAMPOS_FOTO_VALIDOS:
        flash("Documento no válido.")
        return redirect(url_for("index"))

    db = get_db()
    jugador = db.execute("SELECT * FROM jugadores WHERE id = ?", (jugador_id,)).fetchone()
    if not jugador:
        flash("Jugador no encontrado.")
        return redirect(url_for("index"))
    if not jugador_permitido(db, jugador):
        flash("No tienes acceso a ese jugador.")
        return redirect(url_for("index"))

    if jugador["calificado"] and session.get("rol") != "admin":
        flash("Este jugador ya fue calificado y no se puede modificar. Contacta a la Comisión de Calificación.")
        return redirect(url_for("ficha_jugador", jugador_id=jugador_id))

    db.execute(f"UPDATE jugadores SET {campo} = NULL WHERE id = ?", (jugador_id,))
    db.commit()
    flash("Documento eliminado.", "ok")
    return redirect(url_for("ficha_jugador", jugador_id=jugador_id))


@app.route("/jugador/<int:jugador_id>/calificar", methods=["POST"])
@calificacion_required
def calificar_jugador(jugador_id):
    db = get_db()
    jugador = db.execute("SELECT * FROM jugadores WHERE id = ?", (jugador_id,)).fetchone()
    if not jugador:
        flash("Jugador no encontrado.")
        return redirect(url_for("index"))
    nuevo_estado = 0 if jugador["calificado"] else 1
    db.execute("UPDATE jugadores SET calificado = ? WHERE id = ?", (nuevo_estado, jugador_id))
    db.commit()
    flash("Jugador calificado correctamente." if nuevo_estado else "Se quitó la calificación del jugador.", "ok")
    return redirect(request.referrer or url_for("ficha_jugador", jugador_id=jugador_id))


def _font(size, bold=False):
    candidatos = ["arialbd.ttf", "Arial Bold.ttf", "DejaVuSans-Bold.ttf"] if bold else ["arial.ttf", "Arial.ttf", "DejaVuSans.ttf"]
    for nombre in candidatos:
        try:
            return ImageFont.truetype(nombre, size)
        except OSError:
            continue
    return ImageFont.load_default(size=size)


def _texto_centrado(draw, cx, y, texto, font, fill):
    bbox = draw.textbbox((0, 0), texto, font=font)
    ancho_texto = bbox[2] - bbox[0]
    draw.text((cx - ancho_texto / 2, y), texto, font=font, fill=fill)


_LOGO_CACHE = {}


def _logo_thumb(tam):
    """Logo del carnet ya reducido de tamaño, cacheado en memoria para no
    reabrir y decodificar el PNG del disco en cada jugador (se generan
    hasta 35 carnets seguidos por equipo)."""
    if tam not in _LOGO_CACHE:
        logo_path = os.path.join(BASE_DIR, "static", "logo_ldbo.png")
        if os.path.exists(logo_path):
            img = Image.open(logo_path).convert("RGBA")
            img.thumbnail(tam)
            _LOGO_CACHE[tam] = img
        else:
            _LOGO_CACHE[tam] = None
    return _LOGO_CACHE[tam]


def _generar_carnet(jugador):
    # Proporción 93x64mm (tamaño real de la lámina de laminado que se usa para el carnet).
    ancho, alto = 900, 619

    # ---------- FRENTE ----------
    frente = Image.new("RGB", (ancho, alto), "#eef6f0")
    draw = ImageDraw.Draw(frente)

    for y in range(130, alto):
        t = (y - 130) / (alto - 130)
        color = (
            int(240 + (214 - 240) * t),
            int(248 + (232 - 248) * t),
            int(242 + (220 - 242) * t),
        )
        draw.line([(0, y), (ancho, y)], fill=color)

    draw.rectangle([0, 0, ancho, 130], fill="#ffffff")
    draw.rectangle([0, 127, ancho, 130], fill="#14532d")

    f_titulo = _font(32, bold=True)
    f_sub = _font(22, bold=True)
    _texto_centrado(draw, ancho / 2, 18, "LIGA DEPORTIVA OYAMBARILLO", f_titulo, "#14532d")
    _texto_centrado(draw, ancho / 2, 62, "CAMPEONATO OFICIAL 2026", f_sub, "#b45309")

    logo = _logo_thumb((120, 120))
    if logo:
        frente.paste(logo, (18, 5), logo)

    foto_x, foto_y, foto_w, foto_h = 610, 155, 250, 280
    if jugador["foto"]:
        foto_path = os.path.join(FOTOS_DIR, jugador["foto"])
        if os.path.exists(foto_path):
            foto = Image.open(foto_path)
            # Decodifica ya reducido (más rápido y con mucha menos memoria que
            # abrir la foto a su resolución completa solo para achicarla después).
            foto.draft("RGB", (foto_w * 2, foto_h * 2))
            foto = foto.convert("RGB")
            foto = ImageOps.fit(foto, (foto_w, foto_h))
            frente.paste(foto, (foto_x, foto_y))
    draw.rectangle([foto_x, foto_y, foto_x + foto_w, foto_y + foto_h], outline="#14532d", width=4)

    if jugador["numero_camiseta"]:
        f_num = _font(64, bold=True)
        _texto_centrado(draw, foto_x + foto_w / 2, foto_y + foto_h + 14, f"# {jugador['numero_camiseta']}", f_num, "#14532d")

    f_label = _font(24, bold=True)
    f_valor = _font(34, bold=True)

    y0 = 165
    draw.text((30, y0), "EQUIPO:", font=f_label, fill="#b45309")
    draw.text((30, y0 + 32), jugador["equipo"].upper(), font=f_valor, fill="#14532d")

    y1 = y0 + 105
    draw.text((30, y1), "JUGADOR:", font=f_label, fill="#b45309")
    draw.text((30, y1 + 32), jugador["apellidos"].upper(), font=f_valor, fill="#111111")
    draw.text((30, y1 + 72), jugador["nombres"].upper(), font=f_valor, fill="#111111")

    y2 = y1 + 130
    draw.text((30, y2), "C.I.:", font=f_label, fill="#b45309")
    draw.text((115, y2 - 5), jugador["cedula"], font=f_valor, fill="#111111")

    subcat = jugador["subcategoria"] or "Sub 45"
    f_cat = _font(26, bold=True)
    draw.text((30, y2 + 45), subcat.upper(), font=f_cat, fill="#14532d")

    # ---------- REVERSO ----------
    reverso = Image.new("RGB", (ancho, alto), "white")
    rdraw = ImageDraw.Draw(reverso)

    edad = _calcular_edad(jugador["fecha_nacimiento"])
    f_rlabel = _font(30, bold=True)
    f_rvalor = _font(30)

    rdraw.text((50, 70), "Campeonato:", font=f_rlabel, fill="#14532d")
    rdraw.text((300, 70), (jugador["categoria"] or "") + (f" - {subcat}" if subcat == "Juvenil" else ""), font=f_rvalor, fill="black")

    rdraw.text((50, 130), "Edad:", font=f_rlabel, fill="black")
    rdraw.text((300, 130), str(edad) if edad is not None else "-", font=f_rvalor, fill="black")

    rdraw.text((50, 190), "F. Nacimiento:", font=f_rlabel, fill="black")
    rdraw.text((50, 232), jugador["fecha_nacimiento"] or "-", font=f_rvalor, fill="#14532d")

    sello = _logo_thumb((220, 220))
    if sello:
        sello = sello.copy()
        sello_alpha = sello.split()[3].point(lambda p: p * 0.5)
        sello.putalpha(sello_alpha)
        reverso.paste(sello, (ancho - sello.width - 60, 280), sello)

    rdraw.line([(50, 490), (380, 490)], fill="black", width=2)
    f_firma = _font(24, bold=True)
    rdraw.text((50, 498), "PRESIDENTE", font=f_firma, fill="black")

    return frente, reverso


def _carnet_response(jugador):
    frente, reverso = _generar_carnet(jugador)

    lienzo = Image.new("RGB", (frente.width, frente.height * 2 + 20), "#dddddd")
    lienzo.paste(frente, (0, 0))
    lienzo.paste(reverso, (0, frente.height + 20))

    buf = io.BytesIO()
    lienzo.save(buf, format="PNG")
    buf.seek(0)
    nombre_archivo = f"carnet_{jugador['cedula']}.png"
    return send_file(buf, mimetype="image/png", as_attachment=False, download_name=nombre_archivo)


@app.route("/jugador/<int:jugador_id>/carnet")
@login_required
def carnet_jugador(jugador_id):
    db = get_db()
    jugador = db.execute("SELECT * FROM jugadores WHERE id = ?", (jugador_id,)).fetchone()
    if not jugador:
        flash("Jugador no encontrado.")
        return redirect(url_for("index"))
    if not jugador_permitido(db, jugador):
        flash("No tienes acceso a ese jugador.")
        return redirect(url_for("index"))
    return _carnet_response(jugador)


@app.route("/jugador/<int:jugador_id>/carnet.pdf")
@login_required
def carnet_jugador_pdf(jugador_id):
    db = get_db()
    jugador = db.execute("SELECT * FROM jugadores WHERE id = ?", (jugador_id,)).fetchone()
    if not jugador:
        flash("Jugador no encontrado.")
        return redirect(url_for("index"))
    if not jugador_permitido(db, jugador):
        flash("No tienes acceso a ese jugador.")
        return redirect(url_for("index"))

    frente, reverso = _generar_carnet(jugador)
    buf = io.BytesIO()
    frente.save(buf, format="PDF", save_all=True, append_images=[reverso])
    buf.seek(0)
    nombre_archivo = f"carnet_{jugador['cedula']}.pdf"
    return send_file(buf, mimetype="application/pdf", as_attachment=True, download_name=nombre_archivo)


CARNET_ANCHO_MM = 93
CARNET_ALTO_MM = 64
CARNET_PDF_DPI = 120


def _mm_a_px(mm, dpi=CARNET_PDF_DPI):
    return int(round(mm / 25.4 * dpi))


def _armar_paginas_carnets(jugadores, dpi=CARNET_PDF_DPI):
    """Arma páginas A4 con los carnets en cuadrícula. El reverso de cada jugador va a
    la izquierda y el frente a la derecha, uno al lado del otro y separados por una
    línea guía, para que no haya que buscar el reverso en otra página."""
    margen = _mm_a_px(8, dpi)
    espacio = _mm_a_px(4, dpi)
    espacio_interno = _mm_a_px(3, dpi)
    pagina_w, pagina_h = _mm_a_px(210, dpi), _mm_a_px(297, dpi)
    carnet_w, carnet_h = _mm_a_px(CARNET_ANCHO_MM, dpi), _mm_a_px(CARNET_ALTO_MM, dpi)
    unidad_w = carnet_w * 2 + espacio_interno
    unidad_h = carnet_h

    columnas = max(1, (pagina_w - 2 * margen + espacio) // (unidad_w + espacio))
    filas = max(1, (pagina_h - 2 * margen + espacio) // (unidad_h + espacio))
    por_pagina = columnas * filas

    paginas = []
    for inicio in range(0, len(jugadores), por_pagina):
        lote = jugadores[inicio:inicio + por_pagina]
        pagina = Image.new("RGB", (pagina_w, pagina_h), "white")
        draw = ImageDraw.Draw(pagina)
        for i, jugador in enumerate(lote):
            fila, col = divmod(i, columnas)
            x = margen + col * (unidad_w + espacio)
            y = margen + fila * (unidad_h + espacio)
            frente, reverso = _generar_carnet(jugador)
            pagina.paste(reverso.resize((carnet_w, carnet_h), Image.LANCZOS), (x, y))
            linea_x = x + carnet_w + espacio_interno // 2
            draw.line([(linea_x, y), (linea_x, y + carnet_h)], fill="#bbbbbb", width=2)
            pagina.paste(frente.resize((carnet_w, carnet_h), Image.LANCZOS), (x + carnet_w + espacio_interno, y))
        paginas.append(pagina)
    return paginas


@app.route("/carnets")
@imprenta_required
def carnets_modulo():
    db = get_db()
    equipos = db.execute("SELECT * FROM equipos ORDER BY categoria, division, nombre").fetchall()
    equipo_id = request.args.get("equipo_id", "").strip()
    jugadores = []
    equipo_actual = None
    if equipo_id:
        equipo_actual = db.execute("SELECT * FROM equipos WHERE id = ?", (equipo_id,)).fetchone()
        if equipo_actual:
            jugadores = db.execute(
                "SELECT * FROM jugadores WHERE equipo = ? ORDER BY apellidos, nombres",
                (equipo_actual["nombre"],),
            ).fetchall()
    total = len(jugadores)
    impresos = sum(1 for j in jugadores if j["carnet_impreso"])
    recaudado = sum((j["carnet_valor"] or 0) for j in jugadores if j["carnet_impreso"])
    return render_template(
        "carnets_modulo.html",
        equipos=equipos,
        equipo_actual=equipo_actual,
        jugadores=jugadores,
        total=total,
        impresos=impresos,
        recaudado=recaudado,
    )


@app.route("/carnets/<int:jugador_id>/estado", methods=["POST"])
@imprenta_required
def carnets_estado(jugador_id):
    db = get_db()
    jugador = db.execute("SELECT * FROM jugadores WHERE id = ?", (jugador_id,)).fetchone()
    if not jugador:
        flash("Jugador no encontrado.")
        return redirect(url_for("carnets_modulo"))

    impreso = 1 if request.form.get("carnet_impreso") == "on" else 0
    valor = _to_float(request.form.get("carnet_valor", "0"))
    fecha = jugador["carnet_fecha"]
    if impreso and not jugador["carnet_impreso"]:
        fecha = datetime.now().strftime("%Y-%m-%d %H:%M")
    elif not impreso:
        fecha = None

    db.execute(
        "UPDATE jugadores SET carnet_impreso = ?, carnet_valor = ?, carnet_fecha = ? WHERE id = ?",
        (impreso, valor, fecha, jugador_id),
    )
    db.commit()
    flash("Carnet actualizado.", "ok")

    equipo = db.execute("SELECT id FROM equipos WHERE nombre = ?", (jugador["equipo"],)).fetchone()
    return redirect(url_for("carnets_modulo", equipo_id=equipo["id"] if equipo else ""))


@app.route("/carnets/pdf")
@imprenta_required
def carnets_pdf():
    db = get_db()
    equipo_id = request.args.get("equipo_id", "").strip()
    equipo = db.execute("SELECT * FROM equipos WHERE id = ?", (equipo_id,)).fetchone() if equipo_id else None
    if not equipo:
        flash("Selecciona un equipo para generar los carnets.")
        return redirect(url_for("carnets_modulo"))

    jugador_ids = [int(x) for x in request.args.getlist("jugador_id") if x.isdigit()]
    if jugador_ids:
        placeholders = ",".join("?" * len(jugador_ids))
        jugadores = db.execute(
            f"SELECT * FROM jugadores WHERE equipo = ? AND id IN ({placeholders}) ORDER BY apellidos, nombres",
            (equipo["nombre"], *jugador_ids),
        ).fetchall()
    else:
        jugadores = db.execute(
            "SELECT * FROM jugadores WHERE equipo = ? ORDER BY apellidos, nombres",
            (equipo["nombre"],),
        ).fetchall()

    if not jugadores:
        flash("No se encontró ningún jugador seleccionado de ese equipo." if jugador_ids
              else "Ese equipo todavía no tiene jugadores inscritos.")
        return redirect(url_for("carnets_modulo", equipo_id=equipo_id))

    paginas = _armar_paginas_carnets(jugadores)
    buf = io.BytesIO()
    paginas[0].save(
        buf, format="PDF", save_all=True, append_images=paginas[1:], resolution=CARNET_PDF_DPI
    )
    buf.seek(0)
    sufijo = "seleccionados" if jugador_ids else "todos"
    nombre_archivo = f"carnets_{equipo['nombre'].replace(' ', '_')}_{sufijo}.pdf"
    return send_file(buf, mimetype="application/pdf", as_attachment=True, download_name=nombre_archivo)


def _fixture_completo(db, categoria=CATEGORIA_ACTIVA, division=""):
    jornadas_rows = db.execute(
        "SELECT * FROM jornadas WHERE categoria = ? AND division = ? ORDER BY numero, id",
        (categoria, division),
    ).fetchall()
    jornadas = []
    for jr in jornadas_rows:
        partidos = db.execute(
            "SELECT * FROM partidos WHERE jornada_id = ? ORDER BY id", (jr["id"],)
        ).fetchall()
        descansos = db.execute(
            "SELECT * FROM descansos WHERE jornada_id = ? ORDER BY id", (jr["id"],)
        ).fetchall()
        jornadas.append({"jornada": jr, "partidos": partidos, "descansos": descansos})
    return jornadas


@app.route("/comision")
@calificacion_required
def comision_modulo():
    db = get_db()
    categoria, division = _categoria_y_division(
        db, request.args.get("categoria", CATEGORIA_ACTIVA), request.args.get("division", "")
    )
    jugadores = db.execute(
        "SELECT * FROM jugadores WHERE categoria = ? AND division = ? ORDER BY equipo, apellidos, nombres",
        (categoria, division),
    ).fetchall()

    return render_template(
        "comision_modulo.html",
        jugadores=jugadores,
        categoria=categoria,
        categorias_liga=_categorias_liga(db),
        division=division,
        divisiones=_divisiones_de_categoria(db, categoria),
    )


@app.route("/tecnica/jornada/agregar", methods=["POST"])
@tecnica_required
def agregar_jornada():
    db = get_db()
    categoria, division = _categoria_y_division(
        db, request.form.get("categoria", CATEGORIA_ACTIVA), request.form.get("division", "")
    )
    numero = request.form.get("numero", "").strip()
    fecha = request.form.get("fecha", "").strip()
    if not numero.isdigit():
        flash("Ingresa un número de jornada válido.")
        return redirect(url_for("tecnica_modulo", categoria=categoria, division=division))

    locales = request.form.getlist("equipo_local")
    visitantes = request.form.getlist("equipo_visitante")
    horas = request.form.getlist("hora")
    fechas_partido = request.form.getlist("fecha_partido")
    descansos = [e for e in request.form.getlist("descansos") if e]

    partidos_validos = []
    equipos_usados = set()
    for local, visitante, hora, fecha_partido in zip(locales, visitantes, horas, fechas_partido):
        local, visitante, hora, fecha_partido = local.strip(), visitante.strip(), hora.strip(), fecha_partido.strip()
        if not local or not visitante:
            continue
        if local == visitante:
            flash("Un equipo no puede jugar contra sí mismo.")
            return redirect(url_for("tecnica_modulo", categoria=categoria, division=division))
        if local in equipos_usados or visitante in equipos_usados:
            flash("Hay un equipo asignado a más de un partido en la misma jornada.")
            return redirect(url_for("tecnica_modulo", categoria=categoria, division=division))
        equipos_usados.add(local)
        equipos_usados.add(visitante)
        partidos_validos.append((local, visitante, hora, fecha_partido))

    for equipo in descansos:
        if equipo in equipos_usados:
            flash(f"'{equipo}' no puede descansar y jugar en la misma jornada.")
            return redirect(url_for("tecnica_modulo", categoria=categoria, division=division))

    if not partidos_validos and not descansos:
        flash("Agrega al menos un partido o un equipo que descansa.")
        return redirect(url_for("tecnica_modulo", categoria=categoria, division=division))

    returning = " RETURNING id" if USANDO_POSTGRES else ""
    cur = db.execute(
        f"INSERT INTO jornadas (numero, fecha, categoria, division) VALUES (?, ?, ?, ?){returning}",
        (int(numero), fecha or None, categoria, division),
    )
    jornada_id = cur.fetchone()["id"] if USANDO_POSTGRES else cur.lastrowid
    for local, visitante, hora, fecha_partido in partidos_validos:
        db.execute(
            "INSERT INTO partidos (jornada_id, equipo_local, equipo_visitante, hora, fecha) VALUES (?, ?, ?, ?, ?)",
            (jornada_id, local, visitante, hora or None, fecha_partido or None),
        )
    for equipo in descansos:
        db.execute("INSERT INTO descansos (jornada_id, equipo) VALUES (?, ?)", (jornada_id, equipo))
    db.commit()
    flash(f"Jornada {numero} agregada.", "ok")
    return redirect(url_for("tecnica_modulo", categoria=categoria, division=division))


@app.route("/tecnica/jornada/<int:jornada_id>/eliminar", methods=["POST"])
@tecnica_required
def eliminar_jornada(jornada_id):
    db = get_db()
    jornada = db.execute("SELECT * FROM jornadas WHERE id = ?", (jornada_id,)).fetchone()
    categoria = jornada["categoria"] if jornada else CATEGORIA_ACTIVA
    division = jornada["division"] if jornada else ""
    db.execute("DELETE FROM partidos WHERE jornada_id = ?", (jornada_id,))
    db.execute("DELETE FROM descansos WHERE jornada_id = ?", (jornada_id,))
    db.execute("DELETE FROM jornadas WHERE id = ?", (jornada_id,))
    db.commit()
    flash("Jornada eliminada.", "ok")
    return redirect(url_for("tecnica_modulo", categoria=categoria, division=division))


@app.route("/sanciones")
@sanciones_required
def sanciones_modulo():
    db = get_db()
    jugadores = db.execute(
        "SELECT * FROM jugadores ORDER BY categoria, equipo, apellidos, nombres"
    ).fetchall()
    jornadas = db.execute("SELECT * FROM jornadas ORDER BY categoria, numero, id").fetchall()

    tarjetas = db.execute("""
        SELECT t.*, j.nombres AS j_nombres, j.apellidos AS j_apellidos, j.equipo AS j_equipo,
               jo.numero AS jornada_numero
        FROM tarjetas t
        JOIN jugadores j ON j.id = t.jugador_id
        LEFT JOIN jornadas jo ON jo.id = t.jornada_id
        ORDER BY t.fecha DESC, t.id DESC
    """).fetchall()

    acumulado_map = {}
    orden_map = []
    for t in tarjetas:
        key = t["jugador_id"]
        if key not in acumulado_map:
            acumulado_map[key] = {
                "nombres": t["j_nombres"], "apellidos": t["j_apellidos"], "equipo": t["j_equipo"],
                "amarillas": 0, "rojas": 0, "multas": 0.0,
            }
            orden_map.append(key)
        if t["tipo"] == "amarilla":
            acumulado_map[key]["amarillas"] += 1
        else:
            acumulado_map[key]["rojas"] += 1
        acumulado_map[key]["multas"] += t["valor_multa"] or 0

    acumulado = sorted(
        acumulado_map.values(),
        key=lambda x: (x["rojas"], x["amarillas"]),
        reverse=True,
    )

    sanciones = db.execute("""
        SELECT s.*, j.nombres AS j_nombres, j.apellidos AS j_apellidos, j.equipo AS j_equipo,
               jo.numero AS jornada_desde_numero
        FROM sanciones s
        JOIN jugadores j ON j.id = s.jugador_id
        LEFT JOIN jornadas jo ON jo.id = s.jornada_desde_id
        ORDER BY s.fecha DESC, s.id DESC
    """).fetchall()

    jugadores_js = [
        {"id": j["id"], "label": f"{j['apellidos']} {j['nombres']} ({j['cedula']})", "equipo": j["equipo"]}
        for j in jugadores
    ]

    return render_template(
        "sanciones_modulo.html",
        jugadores=jugadores,
        jugadores_js=jugadores_js,
        jornadas=jornadas,
        tarjetas=tarjetas,
        acumulado=acumulado,
        sanciones=sanciones,
    )


@app.route("/sanciones/tarjeta/agregar", methods=["POST"])
@sanciones_required
def agregar_tarjeta():
    db = get_db()
    jugador_id = request.form.get("jugador_id", "").strip()
    tipo = request.form.get("tipo", "").strip()
    jornada_id = request.form.get("jornada_id", "").strip() or None
    valor_multa = _to_float(request.form.get("valor_multa", "0"))
    observacion = request.form.get("observacion", "").strip()

    if not jugador_id.isdigit() or tipo not in ("amarilla", "roja"):
        flash("Selecciona un jugador y el tipo de tarjeta.")
        return redirect(url_for("sanciones_modulo"))

    db.execute(
        "INSERT INTO tarjetas (jugador_id, jornada_id, tipo, valor_multa, observacion, fecha) VALUES (?, ?, ?, ?, ?, ?)",
        (int(jugador_id), jornada_id, tipo, valor_multa, observacion or None, datetime.now().strftime("%Y-%m-%d %H:%M")),
    )
    db.commit()
    flash("Tarjeta registrada.", "ok")
    return redirect(url_for("sanciones_modulo"))


@app.route("/sanciones/tarjeta/<int:tarjeta_id>/eliminar", methods=["POST"])
@sanciones_required
def eliminar_tarjeta(tarjeta_id):
    db = get_db()
    db.execute("DELETE FROM tarjetas WHERE id = ?", (tarjeta_id,))
    db.commit()
    flash("Tarjeta eliminada.", "ok")
    return redirect(url_for("sanciones_modulo"))


@app.route("/sanciones/sancion/agregar", methods=["POST"])
@sanciones_required
def agregar_sancion():
    db = get_db()
    jugador_id = request.form.get("jugador_id", "").strip()
    motivo = request.form.get("motivo", "").strip()
    jornadas_sancionado = request.form.get("jornadas_sancionado", "1").strip()
    jornada_desde_id = request.form.get("jornada_desde_id", "").strip() or None
    valor_multa = _to_float(request.form.get("valor_multa", "0"))

    if not jugador_id.isdigit() or not jornadas_sancionado.isdigit() or int(jornadas_sancionado) < 1:
        flash("Selecciona un jugador y la cantidad de jornadas de sanción.")
        return redirect(url_for("sanciones_modulo"))

    db.execute(
        """INSERT INTO sanciones (jugador_id, motivo, jornadas_sancionado, jornada_desde_id, valor_multa, pagada, fecha)
           VALUES (?, ?, ?, ?, ?, 0, ?)""",
        (int(jugador_id), motivo or None, int(jornadas_sancionado), jornada_desde_id, valor_multa,
         datetime.now().strftime("%Y-%m-%d %H:%M")),
    )
    db.commit()
    flash("Sanción registrada.", "ok")
    return redirect(url_for("sanciones_modulo"))


@app.route("/sanciones/sancion/<int:sancion_id>/eliminar", methods=["POST"])
@sanciones_required
def eliminar_sancion(sancion_id):
    db = get_db()
    db.execute("DELETE FROM sanciones WHERE id = ?", (sancion_id,))
    db.commit()
    flash("Sanción eliminada.", "ok")
    return redirect(url_for("sanciones_modulo"))


@app.route("/sanciones/sancion/<int:sancion_id>/pagada", methods=["POST"])
@sanciones_required
def marcar_sancion_pagada(sancion_id):
    db = get_db()
    sancion = db.execute("SELECT * FROM sanciones WHERE id = ?", (sancion_id,)).fetchone()
    if not sancion:
        flash("Sanción no encontrada.")
        return redirect(url_for("sanciones_modulo"))
    nuevo = 0 if sancion["pagada"] else 1
    db.execute("UPDATE sanciones SET pagada = ? WHERE id = ?", (nuevo, sancion_id))
    db.commit()
    flash("Estado de pago actualizado.", "ok")
    return redirect(url_for("sanciones_modulo"))


def _calcular_tabla_posiciones(db, categoria=CATEGORIA_ACTIVA, division=""):
    equipos = db.execute(
        "SELECT nombre FROM equipos WHERE categoria = ? AND division = ? ORDER BY nombre", (categoria, division)
    ).fetchall()
    tabla = {
        e["nombre"]: {"equipo": e["nombre"], "pj": 0, "g": 0, "e": 0, "p": 0, "gf": 0, "gc": 0, "pts": 0}
        for e in equipos
    }

    partidos = db.execute(
        """SELECT p.* FROM partidos p JOIN jornadas jo ON jo.id = p.jornada_id
           WHERE jo.categoria = ? AND jo.division = ? AND p.jugado = 1
             AND p.goles_local IS NOT NULL AND p.goles_visitante IS NOT NULL""",
        (categoria, division),
    ).fetchall()
    for p in partidos:
        local, visitante = p["equipo_local"], p["equipo_visitante"]
        gl, gv = p["goles_local"], p["goles_visitante"]
        if local not in tabla or visitante not in tabla:
            continue
        tabla[local]["pj"] += 1
        tabla[visitante]["pj"] += 1
        tabla[local]["gf"] += gl
        tabla[local]["gc"] += gv
        tabla[visitante]["gf"] += gv
        tabla[visitante]["gc"] += gl
        if gl > gv:
            tabla[local]["g"] += 1
            tabla[local]["pts"] += 3
            tabla[visitante]["p"] += 1
        elif gl < gv:
            tabla[visitante]["g"] += 1
            tabla[visitante]["pts"] += 3
            tabla[local]["p"] += 1
        else:
            tabla[local]["e"] += 1
            tabla[visitante]["e"] += 1
            tabla[local]["pts"] += 1
            tabla[visitante]["pts"] += 1

    filas = list(tabla.values())
    for f in filas:
        f["dg"] = f["gf"] - f["gc"]
    filas.sort(key=lambda f: (-f["pts"], -f["dg"], -f["gf"]))
    return filas


def _calcular_goleadores(db, categoria=CATEGORIA_ACTIVA, division="", limite=15):
    return db.execute(
        """
        SELECT j.id, j.nombres, j.apellidos, j.equipo, SUM(g.cantidad) AS goles
        FROM goles g
        JOIN jugadores j ON j.id = g.jugador_id
        WHERE j.categoria = ? AND j.division = ?
        GROUP BY j.id, j.nombres, j.apellidos, j.equipo
        ORDER BY goles DESC
        LIMIT ?
        """,
        (categoria, division, limite),
    ).fetchall()


@app.route("/tecnica")
@tecnica_required
def tecnica_modulo():
    db = get_db()
    categoria, division = _categoria_y_division(
        db, request.args.get("categoria", CATEGORIA_ACTIVA), request.args.get("division", "")
    )
    equipos = db.execute(
        "SELECT * FROM equipos WHERE categoria = ? AND division = ? ORDER BY nombre", (categoria, division)
    ).fetchall()
    jornadas = _fixture_completo(db, categoria, division)
    ultima = db.execute(
        "SELECT MAX(numero) m FROM jornadas WHERE categoria = ? AND division = ?", (categoria, division)
    ).fetchone()["m"]
    siguiente_numero = (ultima or 0) + 1
    jugadores = db.execute(
        "SELECT * FROM jugadores WHERE categoria = ? AND division = ? ORDER BY equipo, apellidos, nombres",
        (categoria, division),
    ).fetchall()
    jugadores_js = [
        {"id": j["id"], "label": f"{j['apellidos']} {j['nombres']} ({j['cedula']})", "equipo": j["equipo"]}
        for j in jugadores
    ]
    tabla = _calcular_tabla_posiciones(db, categoria, division)
    goleadores = _calcular_goleadores(db, categoria, division)
    goles_registrados = db.execute(
        """
        SELECT g.*, j.nombres AS j_nombres, j.apellidos AS j_apellidos, j.equipo AS j_equipo
        FROM goles g
        JOIN jugadores j ON j.id = g.jugador_id
        WHERE j.categoria = ? AND j.division = ?
        ORDER BY g.fecha DESC, g.id DESC
        """,
        (categoria, division),
    ).fetchall()

    return render_template(
        "tecnica_modulo.html",
        categoria=categoria,
        categorias_liga=_categorias_liga(db),
        division=division,
        divisiones=_divisiones_de_categoria(db, categoria),
        equipos=equipos,
        siguiente_numero=siguiente_numero,
        jornadas=jornadas,
        jugadores_js=jugadores_js,
        tabla=tabla,
        goleadores=goleadores,
        goles_registrados=goles_registrados,
    )


def _categoria_de_partido(db, partido_id):
    fila = db.execute(
        "SELECT jo.categoria AS categoria, jo.division AS division "
        "FROM partidos p JOIN jornadas jo ON jo.id = p.jornada_id WHERE p.id = ?",
        (partido_id,),
    ).fetchone()
    return (fila["categoria"], fila["division"]) if fila else (CATEGORIA_ACTIVA, "")


@app.route("/tecnica/partido/<int:partido_id>/resultado", methods=["POST"])
@tecnica_required
def registrar_resultado(partido_id):
    db = get_db()
    partido = db.execute("SELECT * FROM partidos WHERE id = ?", (partido_id,)).fetchone()
    categoria, division = _categoria_de_partido(db, partido_id)
    if not partido:
        flash("Partido no encontrado.")
        return redirect(url_for("tecnica_modulo", categoria=categoria, division=division))
    gl = request.form.get("goles_local", "").strip()
    gv = request.form.get("goles_visitante", "").strip()
    if not gl.isdigit() or not gv.isdigit():
        flash("Ingresa un marcador válido (números).")
        return redirect(url_for("tecnica_modulo", categoria=categoria, division=division))
    db.execute(
        "UPDATE partidos SET goles_local = ?, goles_visitante = ?, jugado = 1 WHERE id = ?",
        (int(gl), int(gv), partido_id),
    )
    db.commit()
    flash("Resultado registrado.", "ok")
    return redirect(url_for("tecnica_modulo", categoria=categoria, division=division))


@app.route("/tecnica/partido/<int:partido_id>/quitar_resultado", methods=["POST"])
@tecnica_required
def quitar_resultado(partido_id):
    db = get_db()
    categoria, division = _categoria_de_partido(db, partido_id)
    db.execute(
        "UPDATE partidos SET goles_local = NULL, goles_visitante = NULL, jugado = 0 WHERE id = ?",
        (partido_id,),
    )
    db.commit()
    flash("Resultado eliminado.", "ok")
    return redirect(url_for("tecnica_modulo", categoria=categoria, division=division))


def _redirect_post_partido(partido_id, categoria, division):
    """Vocalía vuelve a su hoja del encuentro; el resto (técnica/admin) vuelve
    al módulo de la Comisión Técnica."""
    if session.get("rol") == "vocalia":
        return redirect(url_for("vocalia_hoja", partido_id=partido_id))
    return redirect(url_for("tecnica_modulo", categoria=categoria, division=division))


@app.route("/tecnica/partido/<int:partido_id>/en_vivo/iniciar", methods=["POST"])
@tecnica_o_vocalia_required
def iniciar_en_vivo(partido_id):
    db = get_db()
    partido = db.execute("SELECT * FROM partidos WHERE id = ?", (partido_id,)).fetchone()
    categoria, division = _categoria_de_partido(db, partido_id)
    if not partido:
        flash("Partido no encontrado.")
        return _redirect_post_partido(partido_id, categoria, division)
    if partido["jugado"]:
        flash("Ese partido ya está finalizado.")
        return _redirect_post_partido(partido_id, categoria, division)
    goles_local = partido["goles_local"] if partido["goles_local"] is not None else 0
    goles_visitante = partido["goles_visitante"] if partido["goles_visitante"] is not None else 0
    db.execute(
        "UPDATE partidos SET en_vivo = 1, goles_local = ?, goles_visitante = ? WHERE id = ?",
        (goles_local, goles_visitante, partido_id),
    )
    db.commit()
    flash("Partido marcado como EN VIVO — ya se ve en la página pública.", "ok")
    return _redirect_post_partido(partido_id, categoria, division)


@app.route("/tecnica/partido/<int:partido_id>/en_vivo/actualizar", methods=["POST"])
@tecnica_o_vocalia_required
def actualizar_marcador_en_vivo(partido_id):
    db = get_db()
    categoria, division = _categoria_de_partido(db, partido_id)
    gl = request.form.get("goles_local", "").strip()
    gv = request.form.get("goles_visitante", "").strip()
    if not gl.isdigit() or not gv.isdigit():
        flash("Ingresa un marcador válido (números).")
        return _redirect_post_partido(partido_id, categoria, division)
    db.execute(
        "UPDATE partidos SET goles_local = ?, goles_visitante = ? WHERE id = ? AND en_vivo = 1",
        (int(gl), int(gv), partido_id),
    )
    db.commit()
    flash("Marcador actualizado.", "ok")
    return _redirect_post_partido(partido_id, categoria, division)


@app.route("/tecnica/partido/<int:partido_id>/en_vivo/finalizar", methods=["POST"])
@tecnica_o_vocalia_required
def finalizar_en_vivo(partido_id):
    db = get_db()
    categoria, division = _categoria_de_partido(db, partido_id)
    db.execute("UPDATE partidos SET jugado = 1, en_vivo = 0 WHERE id = ?", (partido_id,))
    db.commit()
    flash("Partido finalizado.", "ok")
    return _redirect_post_partido(partido_id, categoria, division)


@app.route("/tecnica/partido/<int:partido_id>/en_vivo/cancelar", methods=["POST"])
@tecnica_o_vocalia_required
def cancelar_en_vivo(partido_id):
    db = get_db()
    categoria, division = _categoria_de_partido(db, partido_id)
    db.execute("UPDATE partidos SET en_vivo = 0 WHERE id = ?", (partido_id,))
    db.commit()
    flash("Se quitó el estado EN VIVO de ese partido.", "ok")
    return _redirect_post_partido(partido_id, categoria, division)


def _partidos_en_vivo(db):
    return db.execute(
        """SELECT p.*, jo.numero AS jornada_numero
           FROM partidos p JOIN jornadas jo ON jo.id = p.jornada_id
           WHERE p.en_vivo = 1 ORDER BY p.id DESC"""
    ).fetchall()


# ---------- Vocalía: hoja oficial del encuentro (asistencia, tarjetas, goles, incidentes, árbitro) ----------

@app.route("/vocalia")
@vocalia_required
def vocalia_modulo():
    db = get_db()
    categoria, division = _categoria_y_division(
        db, request.args.get("categoria", CATEGORIA_ACTIVA), request.args.get("division", "")
    )
    jornadas = _fixture_completo(db, categoria, division)
    return render_template(
        "vocalia_modulo.html",
        categoria=categoria, categorias_liga=_categorias_liga(db),
        division=division, divisiones=_divisiones_de_categoria(db, categoria),
        jornadas=jornadas,
    )


@app.route("/vocalia/partido/<int:partido_id>", methods=["GET"])
@vocalia_required
def vocalia_hoja(partido_id):
    db = get_db()
    partido = db.execute("SELECT * FROM partidos WHERE id = ?", (partido_id,)).fetchone()
    if not partido:
        flash("Partido no encontrado.")
        return redirect(url_for("vocalia_modulo"))
    jornada = db.execute("SELECT * FROM jornadas WHERE id = ?", (partido["jornada_id"],)).fetchone()

    jugadores_local = db.execute(
        "SELECT * FROM jugadores WHERE equipo = ? ORDER BY apellidos, nombres", (partido["equipo_local"],)
    ).fetchall()
    jugadores_visitante = db.execute(
        "SELECT * FROM jugadores WHERE equipo = ? ORDER BY apellidos, nombres", (partido["equipo_visitante"],)
    ).fetchall()

    asistencia_rows = db.execute(
        "SELECT jugador_id, participo FROM vocalia_asistencia WHERE partido_id = ?", (partido_id,)
    ).fetchall()
    asistencia = {r["jugador_id"]: r["participo"] for r in asistencia_rows}

    tarjetas_rows = db.execute(
        "SELECT jugador_id, tipo FROM tarjetas WHERE partido_id = ?", (partido_id,)
    ).fetchall()
    amarillas_ids = {r["jugador_id"] for r in tarjetas_rows if r["tipo"] == "amarilla"}
    rojas_ids = {r["jugador_id"] for r in tarjetas_rows if r["tipo"] == "roja"}

    goles_rows = db.execute(
        "SELECT jugador_id, SUM(cantidad) AS total FROM goles WHERE partido_id = ? GROUP BY jugador_id", (partido_id,)
    ).fetchall()
    goles_por_jugador = {r["jugador_id"]: r["total"] for r in goles_rows}

    incidentes = db.execute(
        "SELECT * FROM vocalia_incidentes WHERE partido_id = ? ORDER BY fecha", (partido_id,)
    ).fetchall()

    cambios = db.execute(
        """SELECT c.*, js.apellidos AS sale_apellidos, js.nombres AS sale_nombres, js.numero_camiseta AS sale_numero,
                  je.apellidos AS entra_apellidos, je.nombres AS entra_nombres, je.numero_camiseta AS entra_numero
           FROM vocalia_cambios c
           JOIN jugadores js ON js.id = c.jugador_sale_id
           JOIN jugadores je ON je.id = c.jugador_entra_id
           WHERE c.partido_id = ? ORDER BY c.fecha""",
        (partido_id,),
    ).fetchall()

    jugadores_js = [
        {"id": j["id"], "label": f"{j['apellidos']} {j['nombres']} ({j['numero_camiseta'] or '-'}) — {j['equipo']}", "equipo": j["equipo"]}
        for j in list(jugadores_local) + list(jugadores_visitante)
    ]

    return render_template(
        "vocalia_hoja.html",
        partido=partido, jornada=jornada,
        jugadores_local=jugadores_local, jugadores_visitante=jugadores_visitante,
        asistencia=asistencia, amarillas_ids=amarillas_ids, rojas_ids=rojas_ids,
        goles_por_jugador=goles_por_jugador, incidentes=incidentes, cambios=cambios,
        jugadores_js=jugadores_js,
    )


@app.route("/vocalia/partido/<int:partido_id>/arbitro/guardar", methods=["POST"])
@vocalia_required
def guardar_arbitro(partido_id):
    db = get_db()
    arbitro = request.form.get("arbitro", "").strip()
    db.execute("UPDATE partidos SET arbitro = ? WHERE id = ?", (arbitro or None, partido_id))
    db.commit()
    flash("Árbitro guardado.", "ok")
    return redirect(url_for("vocalia_hoja", partido_id=partido_id))


@app.route("/vocalia/partido/<int:partido_id>/hoja/guardar", methods=["POST"])
@vocalia_required
def guardar_hoja_equipo(partido_id):
    db = get_db()
    partido = db.execute("SELECT * FROM partidos WHERE id = ?", (partido_id,)).fetchone()
    if not partido:
        flash("Partido no encontrado.")
        return redirect(url_for("vocalia_modulo"))

    equipo_nombre = request.form.get("equipo_nombre", "").strip()
    if equipo_nombre not in (partido["equipo_local"], partido["equipo_visitante"]):
        flash("Equipo inválido.")
        return redirect(url_for("vocalia_hoja", partido_id=partido_id))

    ids_validos = {
        j["id"] for j in db.execute("SELECT id FROM jugadores WHERE equipo = ?", (equipo_nombre,)).fetchall()
    }
    todos_ids = [int(x) for x in request.form.getlist("todos_ids") if x.isdigit() and int(x) in ids_validos]
    marcados_participo = {int(x) for x in request.form.getlist("participo") if x.isdigit()}
    marcados_amarilla = {int(x) for x in request.form.getlist("amarilla") if x.isdigit()}
    marcados_roja = {int(x) for x in request.form.getlist("roja") if x.isdigit()}

    ahora = datetime.now().strftime("%Y-%m-%d %H:%M")

    for jugador_id in todos_ids:
        db.execute("DELETE FROM vocalia_asistencia WHERE partido_id = ? AND jugador_id = ?", (partido_id, jugador_id))
        db.execute(
            "INSERT INTO vocalia_asistencia (partido_id, jugador_id, participo) VALUES (?, ?, ?)",
            (partido_id, jugador_id, 1 if jugador_id in marcados_participo else 0),
        )

        db.execute("DELETE FROM tarjetas WHERE partido_id = ? AND jugador_id = ?", (partido_id, jugador_id))
        if jugador_id in marcados_amarilla:
            db.execute(
                "INSERT INTO tarjetas (jugador_id, jornada_id, partido_id, tipo, fecha) VALUES (?, ?, ?, 'amarilla', ?)",
                (jugador_id, partido["jornada_id"], partido_id, ahora),
            )
        if jugador_id in marcados_roja:
            db.execute(
                "INSERT INTO tarjetas (jugador_id, jornada_id, partido_id, tipo, fecha) VALUES (?, ?, ?, 'roja', ?)",
                (jugador_id, partido["jornada_id"], partido_id, ahora),
            )

        db.execute("DELETE FROM goles WHERE partido_id = ? AND jugador_id = ?", (partido_id, jugador_id))
        raw = request.form.get(f"goles_{jugador_id}", "0").strip()
        cantidad = int(raw) if raw.isdigit() else 0
        if cantidad > 0:
            db.execute(
                "INSERT INTO goles (jugador_id, partido_id, cantidad, fecha) VALUES (?, ?, ?, ?)",
                (jugador_id, partido_id, cantidad, ahora),
            )

    db.commit()
    flash(f"Hoja de {equipo_nombre} guardada.", "ok")
    return redirect(url_for("vocalia_hoja", partido_id=partido_id))


@app.route("/vocalia/partido/<int:partido_id>/cambio/agregar", methods=["POST"])
@vocalia_required
def agregar_cambio(partido_id):
    db = get_db()
    partido = db.execute("SELECT * FROM partidos WHERE id = ?", (partido_id,)).fetchone()
    if not partido:
        flash("Partido no encontrado.")
        return redirect(url_for("vocalia_modulo"))
    equipo = request.form.get("equipo", "").strip()
    if equipo not in (partido["equipo_local"], partido["equipo_visitante"]):
        flash("Selecciona un equipo válido.")
        return redirect(url_for("vocalia_hoja", partido_id=partido_id))
    sale_id = request.form.get("jugador_sale_id", "").strip()
    entra_id = request.form.get("jugador_entra_id", "").strip()
    minuto = request.form.get("minuto", "").strip()
    if not sale_id.isdigit() or not entra_id.isdigit():
        flash("Selecciona el jugador que sale y el que entra.")
        return redirect(url_for("vocalia_hoja", partido_id=partido_id))
    ids_validos = {
        j["id"] for j in db.execute("SELECT id FROM jugadores WHERE equipo = ?", (equipo,)).fetchall()
    }
    if int(sale_id) not in ids_validos or int(entra_id) not in ids_validos:
        flash("Los jugadores del cambio deben pertenecer al equipo seleccionado.")
        return redirect(url_for("vocalia_hoja", partido_id=partido_id))
    db.execute(
        "INSERT INTO vocalia_cambios (partido_id, equipo, jugador_sale_id, jugador_entra_id, minuto, fecha) "
        "VALUES (?, ?, ?, ?, ?, ?)",
        (partido_id, equipo, int(sale_id), int(entra_id), minuto or None, datetime.now().strftime("%Y-%m-%d %H:%M")),
    )
    db.commit()
    flash("Cambio registrado.", "ok")
    return redirect(url_for("vocalia_hoja", partido_id=partido_id))


@app.route("/vocalia/cambio/<int:cambio_id>/eliminar", methods=["POST"])
@vocalia_required
def eliminar_cambio(cambio_id):
    db = get_db()
    cambio = db.execute("SELECT partido_id FROM vocalia_cambios WHERE id = ?", (cambio_id,)).fetchone()
    partido_id = cambio["partido_id"] if cambio else None
    db.execute("DELETE FROM vocalia_cambios WHERE id = ?", (cambio_id,))
    db.commit()
    flash("Cambio eliminado.", "ok")
    if partido_id:
        return redirect(url_for("vocalia_hoja", partido_id=partido_id))
    return redirect(url_for("vocalia_modulo"))


@app.route("/vocalia/partido/<int:partido_id>/incidente/agregar", methods=["POST"])
@vocalia_required
def agregar_incidente(partido_id):
    db = get_db()
    descripcion = request.form.get("descripcion", "").strip()
    if not descripcion:
        flash("Escribe una descripción del incidente.")
        return redirect(url_for("vocalia_hoja", partido_id=partido_id))
    db.execute(
        "INSERT INTO vocalia_incidentes (partido_id, descripcion, fecha) VALUES (?, ?, ?)",
        (partido_id, descripcion, datetime.now().strftime("%Y-%m-%d %H:%M")),
    )
    db.commit()
    flash("Incidente registrado.", "ok")
    return redirect(url_for("vocalia_hoja", partido_id=partido_id))


@app.route("/vocalia/incidente/<int:incidente_id>/eliminar", methods=["POST"])
@vocalia_required
def eliminar_incidente(incidente_id):
    db = get_db()
    incidente = db.execute("SELECT partido_id FROM vocalia_incidentes WHERE id = ?", (incidente_id,)).fetchone()
    partido_id = incidente["partido_id"] if incidente else None
    db.execute("DELETE FROM vocalia_incidentes WHERE id = ?", (incidente_id,))
    db.commit()
    flash("Incidente eliminado.", "ok")
    if partido_id:
        return redirect(url_for("vocalia_hoja", partido_id=partido_id))
    return redirect(url_for("vocalia_modulo"))


@app.route("/tecnica/gol/agregar", methods=["POST"])
@tecnica_required
def agregar_gol():
    db = get_db()
    jugador_id = request.form.get("jugador_id", "").strip()
    partido_id = request.form.get("partido_id", "").strip() or None
    cantidad = request.form.get("cantidad", "1").strip()

    if not jugador_id.isdigit() or not cantidad.isdigit() or int(cantidad) < 1:
        flash("Selecciona un jugador y una cantidad de goles válida.")
        return redirect(url_for("tecnica_modulo"))

    jugador = db.execute("SELECT categoria, division FROM jugadores WHERE id = ?", (int(jugador_id),)).fetchone()
    categoria = jugador["categoria"] if jugador else CATEGORIA_ACTIVA
    division = jugador["division"] if jugador else ""

    db.execute(
        "INSERT INTO goles (jugador_id, partido_id, cantidad, fecha) VALUES (?, ?, ?, ?)",
        (int(jugador_id), partido_id, int(cantidad), datetime.now().strftime("%Y-%m-%d %H:%M")),
    )
    db.commit()
    flash("Gol(es) registrado(s).", "ok")
    return redirect(url_for("tecnica_modulo", categoria=categoria, division=division))


@app.route("/tecnica/gol/<int:gol_id>/eliminar", methods=["POST"])
@tecnica_required
def eliminar_gol(gol_id):
    db = get_db()
    gol = db.execute(
        "SELECT j.categoria AS categoria, j.division AS division FROM goles g JOIN jugadores j ON j.id = g.jugador_id WHERE g.id = ?",
        (gol_id,),
    ).fetchone()
    categoria = gol["categoria"] if gol else CATEGORIA_ACTIVA
    division = gol["division"] if gol else ""
    db.execute("DELETE FROM goles WHERE id = ?", (gol_id,))
    db.commit()
    flash("Registro de gol eliminado.", "ok")
    return redirect(url_for("tecnica_modulo", categoria=categoria, division=division))


# ---------- Sección pública: consulta libre, sin necesidad de cuenta ----------

@app.route("/publico")
def publico_inicio():
    db = get_db()
    return render_template("publico_inicio.html", partidos_en_vivo=_partidos_en_vivo(db))


@app.route("/publico/tabla")
def publico_tabla():
    db = get_db()
    categoria, division = _categoria_y_division(
        db, request.args.get("categoria", CATEGORIA_ACTIVA), request.args.get("division", "")
    )
    tabla = _calcular_tabla_posiciones(db, categoria, division)
    goleadores = _calcular_goleadores(db, categoria, division)
    return render_template(
        "publico_tabla.html", tabla=tabla, goleadores=goleadores, categoria=categoria,
        categorias_liga=_categorias_liga(db), division=division, divisiones=_divisiones_de_categoria(db, categoria),
    )


@app.route("/publico/calendario")
def publico_calendario():
    db = get_db()
    categoria, division = _categoria_y_division(
        db, request.args.get("categoria", CATEGORIA_ACTIVA), request.args.get("division", "")
    )
    jornadas = _fixture_completo(db, categoria, division)
    return render_template(
        "publico_calendario.html", jornadas=jornadas, categoria=categoria,
        categorias_liga=_categorias_liga(db), division=division, divisiones=_divisiones_de_categoria(db, categoria),
    )


@app.route("/publico/equipos")
def publico_equipos():
    db = get_db()
    categoria, division = _categoria_y_division(
        db, request.args.get("categoria", CATEGORIA_ACTIVA), request.args.get("division", "")
    )
    equipos = db.execute(
        "SELECT * FROM equipos WHERE categoria = ? AND division = ? ORDER BY nombre", (categoria, division)
    ).fetchall()
    conteos = {}
    for e in equipos:
        conteos[e["nombre"]] = db.execute(
            "SELECT COUNT(*) c FROM jugadores WHERE equipo = ? AND categoria = ?",
            (e["nombre"], categoria),
        ).fetchone()["c"]
    return render_template(
        "publico_equipos.html", equipos=equipos, conteos=conteos, categoria=categoria,
        categorias_liga=_categorias_liga(db), division=division, divisiones=_divisiones_de_categoria(db, categoria),
    )


@app.route("/publico/jugadores")
def publico_jugadores():
    db = get_db()
    categoria, division = _categoria_y_division(
        db, request.args.get("categoria", CATEGORIA_ACTIVA), request.args.get("division", "")
    )
    # Solo datos que pueden verse públicamente: sin cédula.
    jugadores = db.execute(
        """SELECT nombres, apellidos, equipo, categoria, subcategoria, numero_camiseta, calificado, foto
           FROM jugadores WHERE categoria = ? AND division = ? ORDER BY equipo, apellidos, nombres""",
        (categoria, division),
    ).fetchall()
    return render_template(
        "publico_jugadores.html", jugadores=jugadores, categoria=categoria,
        categorias_liga=_categorias_liga(db), division=division, divisiones=_divisiones_de_categoria(db, categoria),
    )


@app.route("/publico/resultados")
def publico_resultados():
    db = get_db()
    categoria, division = _categoria_y_division(
        db, request.args.get("categoria", CATEGORIA_ACTIVA), request.args.get("division", "")
    )
    partidos = db.execute(
        """
        SELECT p.*, jo.numero AS jornada_numero
        FROM partidos p
        JOIN jornadas jo ON jo.id = p.jornada_id
        WHERE p.jugado = 1 AND jo.categoria = ? AND jo.division = ?
        ORDER BY jo.numero DESC, p.id DESC
        """,
        (categoria, division),
    ).fetchall()
    goleadores = _calcular_goleadores(db, categoria, division)
    return render_template(
        "publico_resultados.html", partidos=partidos, goleadores=goleadores, categoria=categoria,
        categorias_liga=_categorias_liga(db), division=division, divisiones=_divisiones_de_categoria(db, categoria),
    )


@app.route("/jugador/<int:jugador_id>/eliminar", methods=["POST"])
@login_required
def eliminar_jugador(jugador_id):
    db = get_db()
    jugador = db.execute("SELECT * FROM jugadores WHERE id = ?", (jugador_id,)).fetchone()
    if not jugador:
        flash("Jugador no encontrado.")
        return redirect(url_for("index"))
    if not jugador_permitido(db, jugador):
        flash("No tienes acceso a ese jugador.")
        return redirect(url_for("index"))
    equipo_id = db.execute("SELECT id FROM equipos WHERE nombre = ?", (jugador["equipo"],)).fetchone()["id"]
    if jugador["calificado"] and session.get("rol") != "admin":
        flash("Este jugador ya fue calificado y no se puede eliminar. Contacta a la Comisión de Calificación.")
        return redirect(url_for("detalle_equipo", equipo_id=equipo_id))
    db.execute("DELETE FROM jugadores WHERE id = ?", (jugador_id,))
    db.commit()
    flash("Jugador eliminado.", "ok")
    return redirect(url_for("detalle_equipo", equipo_id=equipo_id))


try:
    init_db()
except Exception as e:
    # Si la base de datos no responde al arrancar (ej. suspendida por el
    # proveedor), no tumbamos todo el proceso: dejamos que arranque igual
    # para que las páginas que no dependen de la base (y los mensajes de
    # error) se puedan mostrar en vez de un 502 total.
    print(f"ADVERTENCIA: no se pudo inicializar la base de datos al arrancar: {e}")

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5001))
    app.run(host="0.0.0.0", port=port, debug=True)
