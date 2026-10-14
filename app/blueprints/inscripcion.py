"""Comisión de Inscripciones: equipos por campeonato, lista general e importación desde Excel."""
from flask import Blueprint, flash, redirect, render_template, request, url_for
from openpyxl import load_workbook

from config import CATEGORIA_ACTIVA, CUPO_MAXIMO_EQUIPO, SUBCATEGORIAS
from app.database import get_db
from app.decorators import admin_required, login_required
from app.services.exportacion import _exportar_jugadores_excel
from app.services.importacion import (
    _normalizar_cedula_importada, _normalizar_fecha_importada, _normalizar_nombre_equipo,
    _normalizar_numero_camiseta_importado,
)
from app.services.jugadores import _insertar_jugador
from app.services.liga import _categoria_y_division, _categorias_liga, _divisiones_de_categoria

bp = Blueprint("inscripcion", __name__)


@bp.route("/jugadores_liga")
@login_required
def jugadores_liga():
    """Lista de todos los jugadores inscritos en la liga, de todos los
    equipos, visible tanto para el admin como para los delegados -- para
    que cualquiera pueda revisar quién está inscrito y en qué equipo."""
    db = get_db()
    categoria, division = _categoria_y_division(
        db, request.args.get("categoria", CATEGORIA_ACTIVA), request.args.get("division", "")
    )
    jugadores = db.execute(
        "SELECT * FROM jugadores WHERE categoria = ? AND division = ? ORDER BY equipo, apellidos",
        (categoria, division),
    ).fetchall()
    equipos_distintos = sorted(set(j["equipo"] for j in jugadores))
    return render_template(
        "inscripcion/jugadores_liga.html",
        jugadores=jugadores,
        categoria=categoria,
        categorias_liga=_categorias_liga(db),
        division=division,
        divisiones=_divisiones_de_categoria(db, categoria),
        total_equipos=len(equipos_distintos),
    )


@bp.route("/exportar_general")
@admin_required
def exportar_general():
    db = get_db()
    categoria, division = _categoria_y_division(
        db, request.args.get("categoria", CATEGORIA_ACTIVA), request.args.get("division", "")
    )
    jugadores = db.execute(
        "SELECT * FROM jugadores WHERE categoria = ? AND division = ? ORDER BY equipo, apellidos",
        (categoria, division),
    ).fetchall()
    sufijo = f"{categoria}_{division}" if division else categoria
    nombre_archivo = f"jugadores_liga_oyambarillo_{sufijo.replace(' ', '_')}.xlsx"
    return _exportar_jugadores_excel(jugadores, nombre_archivo, incluir_equipo=True)


@bp.route("/admin/importar_jugadores", methods=["GET", "POST"])
@admin_required
def importar_jugadores():
    if request.method == "GET":
        return render_template("inscripcion/importar_jugadores.html")

    db = get_db()
    archivo = request.files.get("archivo")
    if not archivo or not archivo.filename:
        flash("Selecciona un archivo Excel (.xlsx) para importar.")
        return redirect(url_for("inscripcion.importar_jugadores"))
    if not archivo.filename.lower().endswith((".xlsx", ".xlsm")):
        flash("El archivo debe ser un Excel (.xlsx o .xlsm).")
        return redirect(url_for("inscripcion.importar_jugadores"))

    try:
        wb = load_workbook(archivo, data_only=True)
        ws = wb.worksheets[0]
    except Exception:
        flash("No se pudo leer el archivo. Verifica que sea un Excel válido y no esté dañado.")
        return redirect(url_for("inscripcion.importar_jugadores"))

    equipos_por_nombre = {
        _normalizar_nombre_equipo(e["nombre"]): e
        for e in db.execute("SELECT nombre, categoria FROM equipos").fetchall()
    }

    permitir_sin_fecha = request.form.get("permitir_sin_fecha") == "1"

    importados = 0
    omitidas_vacias = 0
    errores = []
    pendientes_fecha = []
    fila_num = 1
    for fila in ws.iter_rows(min_row=2, values_only=True):
        fila_num += 1
        celdas = list(fila) + [None] * 6
        equipo_raw, numero_raw, nombres_raw, apellidos_raw, cedula_raw, fecha_raw = celdas[:6]

        if not any([equipo_raw, numero_raw, nombres_raw, apellidos_raw, cedula_raw, fecha_raw]):
            omitidas_vacias += 1
            continue

        equipo_nombre = str(equipo_raw or "").strip()
        nombres = str(nombres_raw or "").strip()
        apellidos = str(apellidos_raw or "").strip()
        etiqueta = f"Fila {fila_num} ({equipo_nombre or '¿?'} — {nombres} {apellidos})".strip()

        if not equipo_nombre:
            errores.append(f"Fila {fila_num}: falta el nombre del equipo.")
            continue

        equipo = equipos_por_nombre.get(_normalizar_nombre_equipo(equipo_nombre))
        if not equipo:
            errores.append(f"{etiqueta}: el equipo '{equipo_nombre}' no existe en el sistema. Créalo primero en Inscripciones.")
            continue

        if not nombres or not apellidos:
            errores.append(f"{etiqueta}: faltan nombres o apellidos.")
            continue

        cedula = _normalizar_cedula_importada(cedula_raw)
        if not cedula:
            errores.append(f"{etiqueta}: la cédula '{cedula_raw}' no es válida (debe tener 10 dígitos).")
            continue

        fecha_nacimiento = _normalizar_fecha_importada(fecha_raw)
        numero_camiseta = _normalizar_numero_camiseta_importado(numero_raw)
        falta_fecha = not fecha_nacimiento

        jugador_id, error = _insertar_jugador(
            db, equipo["nombre"], cedula, nombres, apellidos,
            fecha_nacimiento, equipo["categoria"], numero_camiseta,
            omitir_validacion_edad=(falta_fecha and permitir_sin_fecha),
        )
        if error:
            errores.append(f"{etiqueta}: {error}")
        else:
            importados += 1
            if falta_fecha:
                pendientes_fecha.append(f"{etiqueta} — equipo {equipo['nombre']}")

    return render_template(
        "inscripcion/importar_jugadores_resultado.html",
        importados=importados,
        omitidas_vacias=omitidas_vacias,
        errores=errores,
        pendientes_fecha=pendientes_fecha,
        total_filas=fila_num - 1,
    )


@bp.route("/inscripcion", methods=["GET"])
@admin_required
def inscripcion():
    db = get_db()
    categoria, division = _categoria_y_division(
        db, request.args.get("categoria", CATEGORIA_ACTIVA), request.args.get("division", "")
    )

    equipos_rows = db.execute(
        "SELECT * FROM equipos WHERE categoria = ? AND division = ? ORDER BY nombre", (categoria, division)
    ).fetchall()

    cupos = {}
    for equipo in equipos_rows:
        c = db.execute(
            "SELECT COUNT(*) c FROM jugadores WHERE equipo = ? AND categoria = ?",
            (equipo["nombre"], categoria),
        ).fetchone()["c"]
        cupos[equipo["nombre"]] = c

    return render_template(
        "inscripcion/inscripcion.html",
        equipos=equipos_rows,
        cupos=cupos,
        cupo_maximo=CUPO_MAXIMO_EQUIPO,
        categoria=categoria,
        categorias_liga=_categorias_liga(db),
        division=division,
        divisiones=_divisiones_de_categoria(db, categoria),
        subcategorias=SUBCATEGORIAS,
    )
